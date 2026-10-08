# ----------------------------------------------------------------------------
# Authorship
#   Designed and developed by Elgin Akin (Pekosz Lab, Johns Hopkins Bloomberg
#   School of Public Health), 2026, with assistance from Claude (Anthropic).
#   The attribution line "Developed by Elgin Akin, 2026" (ATTRIBUTION below)
#   must stay in this file and in both outputs: the build checks the HTML
#   report and the slide deck for it and fails if it is missing.
# ----------------------------------------------------------------------------
"""
Influenza surveillance report for the Pekosz Lab Nextstrain builds.

This report strictly complements the Nextstrain build: it summarises what
went into the pipeline, what came out, and the type/subtype/clade frequencies
of the specimens. "Submitted" always means specimens fed into the Nextstrain
pipeline (rows of the pipeline's input metadata), not every specimen collected
or sequenced.

One file, three ways to use it:

    # 1. Build the report (the pipeline's `surveillance_report` rule does this).
    #    Writes an interactive HTML page, a PowerPoint deck, figures and tables.
    python scripts/report.py

    # 2. Explore it live as a marimo app (date pickers, run explorer).
    marimo run scripts/report.py

    # 3. Edit it as a marimo notebook (add plots, inspect data frames).
    marimo edit scripts/report.py

Run `python scripts/report.py --help` for every option. Options can also be
passed to marimo after `--`, e.g. `marimo run scripts/report.py -- --as-of 2026-03-01`.

What it reads (paths are relative to --repo, the repository root):
    fludb.db                                    every uploaded sequence
    {source_dir}/JHH_metadata.txt               every submitted specimen
    {source_dir}/flusort_JHH_metadata.tsv       flusort type/subtype calls
    {source_dir}/JHH_sequences.fasta            submitted segment sequences (header count only)
    data/{subtype}/{segment}/metadata.tsv       build inputs
    results/{subtype}/{segment}/metadata_merged.tsv    Nextclade QC + clades
    results/{subtype}/{segment}/filtered.tsv    augur filter survivors
    results/{subtype}/{segment}/filter_log.tsv  augur filter reasons (optional)
    results/{subtype}/{segment}/tree.nwk        tips surviving augur refine
    auspice/{subtype}/{segment}.json            final Nextstrain build
    config/exclude.tsv                          manual exclusions
    config/report.yaml                          report title, notes, options

What it writes (to --outdir, default reports/):
    surveillance_report.html                    interactive report (any browser, offline)
    surveillance_report.pptx                    static slide deck
    figures/*.png                               every figure, 2x resolution
    tables/*.tsv                                every table behind the report
    report.tsv, report.xlsx                     sequence-level summary (+ QC sheets)

Adding a plot or table
----------------------
Every plot and table is one *item*: a function that takes the ReportData
object and returns a block (Chart, Table, KPIs, Compare, Text, Tabs, Choice)
carrying its title, caption and export options. Register it by adding it to
its section in `report_layout()`. The HTML page, the slide deck, the
figure/table exports and the marimo app all pick it up from there:

    @app.function
    def item_my_plot(d):
        chart = lambda width: alt.Chart(d.lab).mark_bar().encode(...).properties(width=width)
        return Chart("my_plot", "My plot title", chart, caption="What it shows")

A Chart takes a function of `width` so the same plot can fill the browser
("container") and be exported at a fixed pixel width for slides.

This replaces scripts/build-reports.py and scripts/render-reports.qmd.
Per-run benchmarking reports (machine performance) remain in
scripts/run_report.py.
"""

import marimo

__generated_with = "0.25.1"
app = marimo.App(width="full", app_title="JHH Influenza Surveillance Report")

with app.setup:
    import argparse
    import base64
    import datetime as dt
    import html as htmlmod
    import io
    import json
    import re
    import sqlite3
    import sys
    import traceback
    from dataclasses import dataclass, field
    from pathlib import Path
    from typing import Any, Callable

    import altair as alt
    import marimo as mo
    import numpy as np
    import pandas as pd

    # ------------------------------------------------------------------
    # Pipeline vocabulary. Keep in sync with Snakefile / segments.smk.
    # ------------------------------------------------------------------
    BUILD_SUBTYPES = ["h3n2", "h1n1", "vic"]  # directory names
    SEGMENTS = ["pb2", "pb1", "pa", "ha", "np", "na", "mp", "ns"]
    BUILDS = SEGMENTS + ["genome"]
    SUBTYPE_DIR = {"H3N2": "h3n2", "H1N1": "h1n1", "Victoria": "vic"}
    DIR_SUBTYPE = {v: k for k, v in SUBTYPE_DIR.items()}
    SUBTYPE_LABEL = {"H3N2": "A/H3N2", "H1N1": "A/H1N1pdm09", "Victoria": "B/Victoria"}
    SUBTYPE_ORDER = ["H3N2", "H1N1", "Victoria"]
    TYPE_LABEL = {"InfluenzaA": "Influenza A", "InfluenzaB": "Influenza B"}

    # ------------------------------------------------------------------
    # Colour. Categorical hues are assigned in a fixed order and follow the
    # entity (a subtype is always the same colour in every chart).
    # Palette validated for colour-vision deficiency (adjacent pairs).
    # ------------------------------------------------------------------
    PALETTE = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100",
               "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
    PALETTE_DARK = ["#3987e5", "#d95926", "#199e70", "#c98500",
                    "#d55181", "#008300", "#9085e9", "#e66767"]
    OTHER_COLOR = "#a8a7a2"
    SUBTYPE_COLORS = {"H3N2": PALETTE[0], "H1N1": PALETTE[1], "Victoria": PALETTE[2]}
    TYPE_COLORS = {"Influenza A": PALETTE[6], "Influenza B": PALETTE[2]}
    SEQ_RAMP = ["#cde2fb", "#9ec5f4", "#6da7ec", "#3987e5", "#256abf", "#184f95", "#0d366b"]
    INK, INK_2, MUTED = "#0b0b0b", "#52514e", "#898781"
    GRID, BASELINE, SURFACE = "#e1e0d9", "#c3c2b7", "#fcfcfb"

    # Where each specimen ended up. Ordered from success to earliest loss;
    # colours carry status meaning and are always shown with a text label.
    OUTCOMES = [
        ("In final build", "#0ca30c"),
        ("Clock-filter outlier", "#fab219"),
        ("Removed by QC filters", "#ec835a"),
        ("No HA segment", "#898781"),
        ("Subtype/lineage not assigned", "#c3c2b7"),
        ("No sequence data", "#d03b3b"),
    ]
    OUTCOME_ORDER = [o for o, _ in OUTCOMES]
    OUTCOME_COLORS = dict(OUTCOMES)

    # Per-build status codes (one per specimen x build)
    STATUS_LABEL = {
        "in_build": "In final build",
        "clock": "Clock-filter outlier",
        "export": "Dropped at export",
        "qc": "Failed Nextclade QC / coverage",
        "length": "Shorter than minimum length",
        "excluded": "Manually excluded (exclude.tsv)",
        "no_sequence": "No sequence data",
        "filter_other": "Removed by augur filter (other)",
        "absent": "Segment not sequenced",
    }
    STATUS_SHORT = {
        "in_build": "✓", "clock": "clock", "export": "export", "qc": "QC",
        "length": "short", "excluded": "excl", "no_sequence": "no seq",
        "filter_other": "filter", "absent": "–",
    }
    STATUS_COLORS = {
        "in_build": "#0ca30c", "clock": "#fab219", "export": "#fab219",
        "qc": "#ec835a", "length": "#ec835a", "excluded": "#d03b3b",
        "no_sequence": "#d03b3b", "filter_other": "#ec835a", "absent": "#e1e0d9",
    }
    FILTER_STATUSES = ["qc", "length", "excluded", "no_sequence", "filter_other"]

    # Run-level funnel: why specimens leave, in pipeline order
    REMOVAL_REASONS = ["No sequence data", "Subtype/lineage not assigned", "No HA segment",
                       "Failed Nextclade QC / coverage", "HA too short", "Manually excluded",
                       "Other augur filter", "Clock-filter outlier"]
    TYPE_GROUPS = ["Influenza A", "Influenza B", "Untyped", "No sequence data"]
    SUBTYPE_GROUPS = ["A/H3N2", "A/H1N1pdm09", "B/Victoria", "A, subtype not assigned",
                      "B, lineage not assigned", "Untyped", "No sequence data"]
    GROUP_COLORS = {
        "Influenza A": PALETTE[6], "Influenza B": PALETTE[2],
        "A/H3N2": PALETTE[0], "A/H1N1pdm09": PALETTE[1], "B/Victoria": PALETTE[2],
        "A, subtype not assigned": "#b7b0e6", "B, lineage not assigned": "#a3dcc6",
        "Untyped": "#c3c2b7", "No sequence data": "#898781",
    }
    # Per-build stages (columns of the per-build table)
    BUILD_STAGES = ["In final build", "Clock outlier", "Failed QC / coverage", "Too short",
                    "Excluded", "No sequence", "Other filter"]
    BUILD_STAGE_COLORS = ["#0ca30c", "#fab219", "#ec835a", "#d03b3b", "#4a3aa7", "#898781", "#c3c2b7"]

    # ------------------------------------------------------------------
    # Wording used in both outputs
    # ------------------------------------------------------------------
    ATTRIBUTION = "Developed by Elgin Akin, 2026"   # required in both outputs (checked)
    SUBMITTED_DEFINITION = (
        "“Submitted” means specimens fed into the Nextstrain pipeline (rows of the pipeline's "
        "input metadata), not every specimen collected or sequenced.")
    COMPLEMENT_STATEMENT = "This report strictly complements the Nextstrain build."
    SEASON_MARKER_NOTE = (
        "Dotted lines at October 1 and May 31 are a visual aid only, not true season boundaries "
        "(those depend on case counts).")

    DEFAULT_CONFIG = {
        "title": "Johns Hopkins Hospital Influenza Virus Sequencing Frequency Report",
        "subtitle": "Consensus sequencing data is provided by Dr. Heba Mostafa",
        "author": "Pekosz Lab",
        "confidentiality": (
            "CONFIDENTIAL AND PRIVILEGED INFORMATION. This document contains "
            "confidential and privileged information intended solely for "
            "authorized personnel and designated recipients."
        ),
        "source_url": "https://github.com/Pekosz-Lab/nextstrain/blob/main/scripts/report.py",
        "lab_origin": "mostafa_lab",          # fludb database_origin of JHH specimens
        "plot_start_date": "2023-01-01",      # earliest date shown in frequency plots
        "plot_end_date": None,                # latest date shown (None = no limit)
        "exclude_runs_from_plots": [],        # e.g. runs with ambiguous dates
        "top_n_lineages": 8,                  # clades/subclades shown before "Other"
        "control_pattern": r"^(NTC|NEG|POS|PC|NC|BLANK|WATER|H2O)",
        # Keep these in sync with the augur filter rule in segments.smk.
        "qc_min_coverage": 0.9,
        "qc_pass_status": ["good", "mediocre"],
        "notes": [],                          # [{kind: warning|important|note, text: ...}]
        "pptx_table_rows": 14,                # table rows per slide
    }

    # ------------------------------------------------------------------
    # Chart theme (light). Dark mode is applied in the browser by swapping
    # these hexes for their dark-surface steps.
    # ------------------------------------------------------------------
    FONT = "system-ui, -apple-system, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif"

    @alt.theme.register("pekosz", enable=True)
    def _pekosz_theme():
        return alt.theme.ThemeConfig(
            config={
                "background": "transparent",
                "font": FONT,
                "padding": {"left": 4, "right": 12, "top": 8, "bottom": 4},
                "view": {"stroke": None},
                "axis": {
                    "labelColor": INK_2, "titleColor": INK_2, "labelFontSize": 11,
                    "titleFontSize": 12, "titleFontWeight": 500, "gridColor": GRID,
                    "domainColor": BASELINE, "tickColor": BASELINE, "labelPadding": 4,
                    "titlePadding": 8,
                },
                "axisBand": {"grid": False},
                "legend": {
                    "labelColor": INK_2, "titleColor": INK, "labelFontSize": 11,
                    "titleFontSize": 12, "titleFontWeight": 600, "orient": "top",
                    "symbolType": "square", "symbolSize": 120, "columnPadding": 14,
                    "labelLimit": 220,
                },
                "title": {"color": INK, "fontSize": 14, "fontWeight": 600,
                          "anchor": "start", "subtitleColor": INK_2},
                "header": {"labelColor": INK, "titleColor": INK, "labelFontSize": 12,
                           "labelFontWeight": 600},
                "range": {"category": PALETTE},
                "bar": {"stroke": SURFACE, "strokeWidth": 1},
                "rect": {"stroke": SURFACE, "strokeWidth": 1},
                "text": {"color": INK, "fontSize": 11},
                "line": {"strokeWidth": 2},
                "point": {"size": 64, "filled": True},
                "circle": {"size": 64},
            }
        )

    _ = alt.data_transformers.disable_max_rows()


# ======================================================================
# Report building blocks
# ======================================================================


@app.class_definition
@dataclass
class Text:
    """Markdown-ish text. kind: text | note | warning | important | error."""
    body: str
    kind: str = "text"
    pptx: bool = False


@app.class_definition
@dataclass
class KPIs:
    """A row of headline numbers: [(label, value, caption), ...]."""
    items: list
    title: str = "At a glance"
    pptx: bool = True


@app.class_definition
@dataclass
class Compare:
    """Side-by-side panels of headline numbers:
    [(heading, subheading, [(label, value, caption), ...]), ...]."""
    panels: list
    title: str = "At a glance"
    caption: str = ""
    pptx: bool = True


@app.class_definition
@dataclass
class Choice:
    """One of several views picked from a menu (HTML) – e.g. one per season.
    options: [(label, [blocks]), ...]; the first option is the default. In the
    deck the first option is shown in place and the others go to the
    appendix, one slide each."""
    label: str
    options: list
    appendix_title: str = "Appendix – {label}"


@app.class_definition
@dataclass
class Chart:
    """An Altair chart. `build(width)` returns the chart; width is either
    "container" (fill the page) or an int (pixels, used for PNG/slides).
    `static` optionally builds a different chart for slides/PNG (e.g. one
    without dropdown filters)."""
    key: str
    title: str
    build: Callable
    caption: str = ""
    static: Callable | None = None
    min_width: int = 0          # px; wider charts scroll sideways on phones
    pptx: bool = True


@app.class_definition
@dataclass
class Table:
    """A table. interactive=True gives search/sort/filter/export in HTML."""
    key: str
    title: str
    df: Any
    caption: str = ""
    interactive: bool = True
    row_styles: list | None = None   # per row: None | "group" | "indent"
    pptx: bool = True
    pptx_columns: list | None = None
    pptx_max_rows: int | None = None
    page_length: int = 10
    html_df: Any = None          # optional HTML-formatted version for the web page


@app.class_definition
@dataclass
class Tabs:
    """Alternative views of the same thing: [(label, [blocks]), ...]."""
    tabs: list


@app.class_definition
@dataclass
class Section:
    key: str
    title: str
    blocks: list
    intro: str = ""


@app.class_definition
@dataclass
class ReportData:
    cfg: dict
    meta: dict
    seqs: Any          # one row per fludb sequence (all origins) + clades
    lab: Any           # seqs restricted to JHH specimens
    samples: Any       # one row per submitted specimen with its QC journey
    build_qc: Any      # one row per specimen x build
    builds: Any        # one row per build: stage counts
    runs: Any          # one row per sequencing run: QC summary
    run_order: list    # runs sorted by first specimen date
    run_dates: dict    # run -> first specimen date


# ======================================================================
# Small helpers
# ======================================================================


@app.function
def fmt_n(n):
    return f"{int(n):,}" if pd.notna(n) else "–"


@app.function
def fmt_pct(num, den, digits=1):
    return f"{100 * num / den:.{digits}f}%" if den else "–"


@app.function
def fmt_n_pct(num, den, digits=1):
    """'1,234 (56.7%)' – the convention used in the summary tables."""
    pct = round(100 * num / den, digits) if den else 0
    return f"{int(num):,} ({pct:g}%)"


# ---------- Season logic (the single definition used everywhere) ----------
# A Northern Hemisphere season runs from October 1 of year Y to about May 31
# of year Y+1. The current season starts October 1 of this year from October
# onward, otherwise October 1 of last year, so June-September still show the
# season that just ended. For counting, specimens dated June-September belong
# to the season that started the previous October (seasons are Oct 1 - Sep 30
# bins); May 31 is only drawn as a visual marker.


@app.function
def season_start_for(date):
    """October 1 that starts the season containing `date`."""
    date = pd.Timestamp(date)
    return pd.Timestamp(year=date.year if date.month >= 10 else date.year - 1, month=10, day=1)


@app.function
def season_end(start):
    """Exclusive end of the counting window of a season (next October 1)."""
    return pd.Timestamp(year=start.year + 1, month=10, day=1)


@app.function
def season_label(start):
    return f"{start.year}-{start.year + 1}"


@app.function
def season_of(date):
    """Season label of a date, e.g. 2024-10-03 -> '2024-2025'."""
    return None if pd.isna(date) else season_label(season_start_for(date))


@app.function
def season_markers(dmin, dmax):
    """(date, label) for every Oct 1 and May 31 between two dates."""
    out = []
    if pd.isna(dmin) or pd.isna(dmax):
        return out
    for y in range(dmin.year - 1, dmax.year + 2):
        for m, day, name in ((5, 31, "May 31"), (10, 1, "Oct 1")):
            t = pd.Timestamp(year=y, month=m, day=day)
            if dmin < t <= dmax:
                out.append((t, f"{name} ’{str(y)[2:]}"))
    return sorted(out)


@app.function
def natural_key(s):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(s))]


@app.function
def read_tsv(path, **kw):
    path = Path(path)
    if not path.exists() or path.stat().st_size == 0:
        return None
    try:
        return pd.read_csv(path, sep="\t", dtype=str, keep_default_na=False, na_values=[""], **kw)
    except pd.errors.EmptyDataError:
        return None


@app.function
def newick_tips(path):
    """Tip names of a Newick tree (internal node labels are skipped)."""
    path = Path(path)
    if not path.exists():
        return None
    text = path.read_text()
    return set(m.strip().strip("'\"") for m in re.findall(r"[(,]([^(),:;]+)", text))


@app.function
def auspice_tips(path):
    """Tip names of an Auspice v2 JSON (iterative, deep trees are fine)."""
    path = Path(path)
    if not path.exists():
        return None
    tree = json.loads(path.read_text())["tree"]
    tips, stack = set(), [tree]
    while stack:
        node = stack.pop()
        children = node.get("children")
        if children:
            stack.extend(children)
        else:
            tips.add(node["name"])
    return tips


@app.function
def fasta_header_ids(path):
    """sequence_ID -> number of segment sequences, from '>ID_segment' headers."""
    path = Path(path)
    if not path.exists():
        return None
    counts = {}
    with open(path) as fh:
        for line in fh:
            if line.startswith(">"):
                sid = line[1:].strip().split()[0].rsplit("_", 1)[0]
                counts[sid] = counts.get(sid, 0) + 1
    return counts


@app.function
def slug(text):
    return re.sub(r"[^a-z0-9]+", "_", str(text).lower()).strip("_")


# ======================================================================
# Configuration and command line
# ======================================================================


@app.function
def parse_args(argv=None):
    p = argparse.ArgumentParser(
        prog="report.py",
        description="Build the JHH influenza surveillance report (HTML + PPTX).",
    )
    p.add_argument("--repo", default=".", help="repository root (default: current directory)")
    p.add_argument("--config", default="config/report.yaml",
                   help="report options YAML (title, notes, plot window; default: config/report.yaml)")
    p.add_argument("--pipeline-config", default=None,
                   help="Snakemake --configfile used for the run (e.g. config/tutorial.yaml); "
                        "used to locate source_dir / jhh_metadata")
    p.add_argument("--source-dir", default=None, help="override source directory (default: source)")
    p.add_argument("--jhh-metadata", default=None, help="submitted specimen metadata (default: {source_dir}/JHH_metadata.txt)")
    p.add_argument("--jhh-sequences", default=None, help="submitted sequences FASTA (default: {source_dir}/JHH_sequences.fasta)")
    p.add_argument("--flusort-metadata", default=None,
                   help="flusort typing output (default: {source_dir}/flusort_JHH_metadata.tsv)")
    p.add_argument("--db", default="fludb.db", help="fludb SQLite database")
    p.add_argument("--outdir", default="reports", help="output directory (default: reports)")
    p.add_argument("--run-id", default=None,
                   help="pipeline RUN_ID (YYYYMMDDTHHMMSS); default: latest in .run/run_logs")
    p.add_argument("--as-of", default=None,
                   help="report date (YYYY-MM-DD); sets the current season. Default: today")
    p.add_argument("--formats", default="html,pptx,png,tsv,xlsx",
                   help="comma-separated outputs to write (default: html,pptx,png,tsv,xlsx)")
    p.add_argument("--strict", action="store_true",
                   help="fail if any section errors (default: show the error in the report)")
    args, _unknown = p.parse_known_args(argv)
    return args


@app.function
def load_yaml(path):
    path = Path(path)
    if not path.exists():
        return {}
    import yaml
    return yaml.safe_load(path.read_text()) or {}


@app.function
def resolve_inputs(args):
    """Turn CLI args + configs into concrete paths."""
    repo = Path(args.repo).resolve()
    cfg = dict(DEFAULT_CONFIG)
    cfg.update({k: v for k, v in load_yaml(repo / args.config).items() if v is not None or k in cfg})
    pipe = load_yaml(repo / args.pipeline_config) if args.pipeline_config else {}
    source_dir = args.source_dir or pipe.get("source_dir", "source")
    paths = {
        "repo": repo,
        "db": repo / args.db,
        "jhh_metadata": repo / (args.jhh_metadata or pipe.get("jhh_metadata", f"{source_dir}/JHH_metadata.txt")),
        "jhh_sequences": repo / (args.jhh_sequences or pipe.get("jhh_sequences", f"{source_dir}/JHH_sequences.fasta")),
        "flusort_metadata": repo / (args.flusort_metadata or f"{source_dir}/flusort_JHH_metadata.tsv"),
        "exclude": repo / "config/exclude.tsv",
        "run_logs": repo / ".run/run_logs",
        "outdir": (repo / args.outdir) if not Path(args.outdir).is_absolute() else Path(args.outdir),
    }
    as_of = pd.Timestamp(args.as_of) if args.as_of else None
    return cfg, paths, as_of


@app.function
def prepared_by():
    """Who ran the pipeline: `nextstrain whoami`, else the OS user, else the
    host name, else "Unknown". Never raises."""
    try:
        import subprocess
        out = subprocess.run(["nextstrain", "whoami"], capture_output=True, text=True, timeout=8)
        if out.returncode == 0:
            lines = [l.strip() for l in out.stdout.splitlines()
                     if l.strip() and not l.strip().startswith("#")]
            if lines and "not logged in" not in out.stdout.lower():
                return lines[-1]
    except Exception:
        pass
    try:
        import getpass
        user = getpass.getuser()
        if user:
            return user
    except Exception:
        pass
    try:
        import socket
        host = socket.gethostname()
        if host:
            return host
    except Exception:
        pass
    return "Unknown"


@app.function
def read_provenance(run_logs, run_id=None):
    """Header facts from .run/run_logs/{RUN_ID}_provenance.txt, if present."""
    run_logs = Path(run_logs)
    files = sorted(run_logs.glob("*_provenance.txt")) if run_logs.exists() else []
    if run_id:
        files = [f for f in files if f.name.startswith(run_id)]
    if not files:
        return {"run_id": run_id}
    f = files[-1]
    info = {"run_id": f.name.split("_provenance")[0]}
    lines = f.read_text().splitlines()
    for i, line in enumerate(lines):
        m = re.match(r"^(\w[\w ]*?):\s+(.*)$", line)
        if m and m.group(1) in ("started", "started_utc", "finished_utc", "status", "host"):
            info[m.group(1)] = m.group(2).strip()
        if line.strip() == "== git ==" and i + 1 < len(lines):
            info["commit"] = lines[i + 1].strip()[:10]
        m = re.match(r"^(augur|nextclade|snakemake):\s+(.*)$", line)
        if m:
            info[m.group(1)] = m.group(2).replace(m.group(1), "").strip()
    return info


# ======================================================================
# Data collection
# ======================================================================


@app.function
def load_fludb(db_path):
    """One row per sequence in fludb with segment completeness (the old
    build-reports.py table)."""
    seg_flags = " + ".join(f"(CASE WHEN {s} IS NOT NULL THEN 1 ELSE 0 END)" for s in SEGMENTS)
    seg_list = " || ".join(f"(CASE WHEN {s} IS NOT NULL THEN '{s};' ELSE '' END)" for s in SEGMENTS)
    query = f"""
        SELECT sequence_ID, sample_ID, type, subtype, date, passage_history, study_id,
               sequencing_run, location, database_origin,
               ({seg_flags}) AS segment_count,
               CASE WHEN ha IS NOT NULL THEN 'yes' ELSE 'no' END AS has_ha,
               CASE WHEN ha IS NOT NULL AND na IS NOT NULL THEN 'yes' ELSE 'no' END AS has_ha_and_na,
               RTRIM({seg_list}, ';') AS segments_present
        FROM influenza_genomes
    """
    with sqlite3.connect(db_path) as conn:
        return pd.read_sql_query(query, conn)


@app.function
def load_clades(repo):
    """HA clade calls per sequence from each subtype's HA build metadata."""
    frames = []
    for st in BUILD_SUBTYPES:
        for name in ("metadata_merged.tsv", "metadata.tsv"):
            df = read_tsv(repo / f"results/{st}/ha/{name}")
            if df is not None:
                break
        if df is None:
            continue
        cols = [c for c in ["sequence_ID", "clade", "subclade", "legacy-clade"] if c in df.columns]
        frames.append(df[cols])
    if not frames:
        return pd.DataFrame(columns=["sequence_ID", "clade", "subclade", "legacy-clade"])
    out = pd.concat(frames, ignore_index=True).drop_duplicates("sequence_ID")
    return out.rename(columns={"legacy-clade": "legacy_clade"})


@app.function
def collect_build(repo, st, seg, exclude_ids, cfg):
    """Follow every input strain of one build through the pipeline.

    Returns one row per strain in data/{st}/{seg}/metadata.tsv with the
    stage at which it left the build (or 'in_build')."""
    meta = read_tsv(repo / f"data/{st}/{seg}/metadata.tsv")
    if meta is None:
        return None
    meta = meta.copy()
    meta["strain"] = meta["sample_ID"]
    # Nextclade QC (segments: metadata_merged.tsv; genome: HA nextclade join)
    if seg == "genome":
        qc = read_tsv(repo / f"results/{st}/genome/nextclade.tsv")
        if qc is not None:
            qc = qc.rename(columns={"qc.overallStatus": "qc_overallStatus",
                                    "qc.overallScore": "qc_overallScore"})
    else:
        qc = read_tsv(repo / f"results/{st}/{seg}/metadata_merged.tsv")
    qc_cols = ["qc_overallStatus", "qc_overallScore", "coverage"]
    if qc is not None and "sample_ID" in qc.columns:
        keep = ["sample_ID"] + [c for c in qc_cols if c in qc.columns]
        meta = meta.merge(qc[keep].drop_duplicates("sample_ID"), on="sample_ID", how="left")
    for c in qc_cols:
        if c not in meta.columns:
            meta[c] = np.nan
    filtered = read_tsv(repo / f"results/{st}/{seg}/filtered.tsv")
    flog = read_tsv(repo / f"results/{st}/{seg}/filter_log.tsv")
    tree = newick_tips(repo / f"results/{st}/{seg}/tree.nwk")
    final = auspice_tips(repo / f"auspice/{st}/{seg}.json")
    passed = set(filtered["sample_ID"]) if filtered is not None and "sample_ID" in filtered else None

    log_reason = {}
    if flog is not None:
        for strain, filt in zip(flog["strain"], flog["filter"]):
            log_reason.setdefault(strain, filt)

    sid_mismatch = dict(zip(meta["strain"], meta["sequence_ID"] != meta["sample_ID"]))
    cov = pd.to_numeric(meta["coverage"], errors="coerce")
    status_ok = meta["qc_overallStatus"].isin(cfg["qc_pass_status"])
    query_fail = ~((cov >= cfg["qc_min_coverage"]) & status_ok)

    def filter_reason(i, strain):
        f = log_reason.get(strain)
        if f:
            if "query" in f:
                return "qc"
            if "min_length" in f:
                return "length"
            if "exclude" in f:
                return "excluded"
            if "sequence" in f:
                return "no_sequence"
            return "filter_other"
        # No filter log (older runs): infer from the filter's inputs.
        if strain in exclude_ids:
            return "excluded"
        if seg == "genome" and sid_mismatch.get(strain):
            return "no_sequence"   # genome FASTA is keyed by sequence_ID, metadata by sample_ID
        if seg != "genome" and query_fail.iloc[i]:
            return "qc"
        return "length"

    status, reason = [], []
    for i, strain in enumerate(meta["strain"]):
        if passed is None:
            s = "unknown"
        elif strain not in passed:
            s = filter_reason(i, strain)
        elif tree is not None and strain not in tree:
            s = "clock"
        elif final is not None and strain not in final:
            s = "export"
        else:
            s = "in_build"
        status.append(s)
        detail = STATUS_LABEL.get(s, s)
        if s == "qc":
            bits = []
            if pd.isna(cov.iloc[i]):
                bits.append("no Nextclade result")
            else:
                if cov.iloc[i] < cfg["qc_min_coverage"]:
                    bits.append(f"coverage {cov.iloc[i]:.2f} < {cfg['qc_min_coverage']}")
                st_val = meta["qc_overallStatus"].iloc[i]
                if pd.notna(st_val) and st_val not in cfg["qc_pass_status"]:
                    bits.append(f"QC status '{st_val}'")
            detail = "Failed Nextclade QC: " + ", ".join(bits) if bits else detail
        reason.append(detail)
    meta["status"] = status
    meta["reason"] = reason
    meta["build"] = f"{st}/{seg}"
    meta["build_subtype"] = st
    meta["segment"] = seg
    meta["filter_log"] = flog is not None
    keep = ["build", "build_subtype", "segment", "strain", "sequence_ID", "sample_ID",
            "sequencing_run", "date", "database_origin", "status", "reason",
            "qc_overallStatus", "qc_overallScore", "coverage", "filter_log"]
    return meta[[c for c in keep if c in meta.columns]]


@app.function
def collect(args):
    """Read every input and assemble ReportData."""
    cfg, paths, as_of = resolve_inputs(args)
    repo = paths["repo"]
    warnings = []

    # --- sequences in fludb (+ clade calls) -----------------------------
    if not paths["db"].exists():
        raise FileNotFoundError(f"fludb not found: {paths['db']} (run the pipeline first)")
    seqs = load_fludb(paths["db"]).merge(load_clades(repo), on="sequence_ID", how="left")
    seqs["date"] = pd.to_datetime(seqs["date"], errors="coerce")
    seqs["season"] = seqs["date"].apply(season_of)
    seqs["virus"] = seqs["subtype"].map(SUBTYPE_LABEL)
    seqs["type_label"] = seqs["type"].map(TYPE_LABEL).fillna("Untyped")
    lab = seqs[seqs["database_origin"] == cfg["lab_origin"]].copy()

    # --- every build, every strain ---------------------------------------
    exclude_ids = set()
    ex = paths["exclude"]
    if ex.exists():
        exclude_ids = {l.split("\t")[0].strip() for l in ex.read_text().splitlines()
                       if l.strip() and not l.startswith("#")}
    frames = []
    for st in BUILD_SUBTYPES:
        for seg in BUILDS:
            b = collect_build(repo, st, seg, exclude_ids, cfg)
            if b is not None:
                frames.append(b)
    build_qc = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(
        columns=["build", "build_subtype", "segment", "strain", "sequence_ID", "status"])
    if not frames:
        warnings.append("No build inputs found under data/ – QC sections are empty.")
    elif not build_qc["filter_log"].any():
        warnings.append("augur filter logs (`results/<subtype>/<segment>/filter_log.tsv`) were not found; "
                        "filter reasons were inferred from the filter's inputs.")

    # --- per-build stage counts ------------------------------------------
    rows = []
    for (st, seg), g in build_qc.groupby(["build_subtype", "segment"], sort=False):
        c = g["status"].value_counts()
        lab_g = g[g.get("database_origin", pd.Series(index=g.index, dtype=str)) == cfg["lab_origin"]]
        rows.append({
            "Subtype": SUBTYPE_LABEL[DIR_SUBTYPE[st]], "subtype_dir": st, "Build": seg.upper() if seg != "genome" else "Genome",
            "segment": seg, "Input": len(g), "JHH input": len(lab_g),
            "Failed QC / coverage": int(c.get("qc", 0)),
            "Too short": int(c.get("length", 0)),
            "Excluded": int(c.get("excluded", 0)),
            "No sequence": int(c.get("no_sequence", 0)),
            "Other filter": int(c.get("filter_other", 0)),
            "Clock outlier": int(c.get("clock", 0) + c.get("export", 0)),
            "In final build": int(c.get("in_build", 0)),
        })
    builds = pd.DataFrame(rows)
    if len(builds):
        builds["Retained"] = builds["In final build"] / builds["Input"]

    # --- submitted specimens and their journey ---------------------------
    submitted = read_tsv(paths["jhh_metadata"])
    flusort = read_tsv(paths["flusort_metadata"])
    if submitted is None:
        warnings.append(f"Submitted metadata not found ({paths['jhh_metadata'].name}); "
                        "using the JHH sequences in fludb as the starting point.")
        submitted = lab[["sequence_ID", "sample_ID", "sequencing_run", "date"]].astype(str)
    submitted = submitted.drop_duplicates("sequence_ID").copy()
    if flusort is not None and "subtype" in flusort.columns:
        submitted = submitted.merge(
            flusort[["sequence_ID", "type", "subtype"]].drop_duplicates("sequence_ID"),
            on="sequence_ID", how="left")
    lab_idx = lab.set_index("sequence_ID")
    for col in ["type", "subtype"]:
        fallback = submitted["sequence_ID"].map(lab_idx[col])
        submitted[col] = submitted[col].fillna(fallback) if col in submitted else fallback
    submitted["has_sequence"] = submitted["sequence_ID"].isin(lab["sequence_ID"])
    seg_counts = fasta_header_ids(paths["jhh_sequences"])
    submitted["segments_submitted"] = (submitted["sequence_ID"].map(seg_counts).fillna(0).astype(int)
                                       if seg_counts is not None else np.nan)
    submitted["segment_count"] = submitted["sequence_ID"].map(lab_idx["segment_count"]).fillna(0).astype(int)
    submitted["date"] = pd.to_datetime(submitted["date"], errors="coerce")
    submitted["is_control"] = submitted["sample_ID"].fillna("").str.match(cfg["control_pattern"], case=False)
    submitted["type_label"] = submitted["type"].map(TYPE_LABEL).fillna("Untyped")
    submitted["build_subtype"] = submitted["subtype"].map(SUBTYPE_DIR)
    for col in ["subclade", "clade", "legacy_clade"]:
        if col in lab_idx:
            submitted[col] = submitted["sequence_ID"].map(lab_idx[col])

    # status per build, wide (one column per segment build)
    lab_builds = build_qc[build_qc["sequence_ID"].isin(submitted["sequence_ID"])]
    wide = lab_builds.pivot_table(index="sequence_ID", columns="segment", values="status",
                                  aggfunc="first")
    for seg in BUILDS:
        col = wide[seg] if seg in wide else pd.Series(dtype=str)
        submitted[f"{seg}_status"] = submitted["sequence_ID"].map(col)
        eligible = submitted["build_subtype"].notna() & submitted["has_sequence"]
        submitted.loc[eligible & submitted[f"{seg}_status"].isna(), f"{seg}_status"] = "absent"
    seg_cols = [f"{s}_status" for s in SEGMENTS]
    submitted["segments_in_builds"] = (submitted[seg_cols] == "in_build").sum(axis=1)
    reasons = lab_builds[lab_builds["segment"] == "ha"].set_index("sequence_ID")["reason"]
    submitted["ha_reason"] = submitted["sequence_ID"].map(reasons)

    def outcome(r):
        if not r["has_sequence"]:
            return "No sequence data"
        if pd.isna(r["build_subtype"]):
            return "Subtype/lineage not assigned"
        s = r["ha_status"]
        if s == "in_build":
            return "In final build"
        if s in ("clock", "export"):
            return "Clock-filter outlier"
        if s in FILTER_STATUSES:
            return "Removed by QC filters"
        return "No HA segment"

    def detail(r):
        if r["is_control"]:
            return "Control sample"
        if not r["has_sequence"]:
            return "No sequence in FASTA / not uploaded to fludb"
        if pd.isna(r["build_subtype"]):
            st = r["subtype"]
            if pd.isna(r["type"]):
                return "flusort could not type the specimen"
            return f"flusort: {TYPE_LABEL.get(r['type'], r['type'])}, subtype '{st}' (incomplete)"
        if r["ha_status"] == "absent":
            return "HA segment not sequenced"
        return r["ha_reason"] if pd.notna(r["ha_reason"]) else STATUS_LABEL.get(r["ha_status"], "")

    submitted["outcome"] = submitted.apply(outcome, axis=1)
    submitted["outcome_detail"] = submitted.apply(detail, axis=1)
    submitted["season"] = submitted["date"].apply(season_of)

    # Segment completeness tiers (segments uploaded to fludb; mutually exclusive)
    segs = submitted["sequence_ID"].map(lab_idx["segments_present"]).fillna("")
    has_ha = segs.str.contains(r"\bha\b")
    has_na = segs.str.contains(r"\bna\b")
    submitted["segment_tier"] = np.select(
        [submitted["segment_count"] == 8, has_ha & has_na, has_ha & ~has_na],
        ["Whole genome", "HA + NA", "HA only"], default="")

    # Breakdown groups used by the run-level funnel
    def type_group(r):
        if not r["has_sequence"]:
            return "No sequence data"
        return TYPE_LABEL.get(r["type"], "Untyped")

    def subtype_group(r):
        if not r["has_sequence"]:
            return "No sequence data"
        if r["subtype"] in SUBTYPE_LABEL and pd.notna(r["build_subtype"]):
            return SUBTYPE_LABEL[r["subtype"]]
        if r["type"] == "InfluenzaA":
            return "A, subtype not assigned"
        if r["type"] == "InfluenzaB":
            return "B, lineage not assigned"
        return "Untyped"

    submitted["type_group"] = submitted.apply(type_group, axis=1)
    submitted["subtype_group"] = submitted.apply(subtype_group, axis=1)
    samples = submitted

    # --- per sequencing run ----------------------------------------------
    run_dates = (samples[~samples["is_control"]].groupby("sequencing_run")["date"]
                 .agg(["min", "max"]))
    run_order = sorted(run_dates.index, key=lambda r: (run_dates.loc[r, "min"], natural_key(r)))
    rows = []
    for run in run_order:
        g = samples[(samples["sequencing_run"] == run) & ~samples["is_control"]]
        o = g["outcome"].value_counts()
        ha = g["ha_status"].value_counts()
        tier = g["segment_tier"].value_counts()
        n = len(g)
        first = run_dates.loc[run, "min"]
        rows.append({
            "Run": run,
            "Date": first.date() if pd.notna(first) else None,
            "Season": season_of(first),
            "Submitted": n,
            "With sequence": int(g["has_sequence"].sum()),
            "Influenza A": int((g["type"] == "InfluenzaA").sum()),
            "Influenza B": int((g["type"] == "InfluenzaB").sum()),
            "A/H3N2": int((g["subtype"] == "H3N2").sum()),
            "A/H1N1pdm09": int((g["subtype"] == "H1N1").sum()),
            "B/Victoria": int((g["subtype"] == "Victoria").sum()),
            "Subtype/lineage not assigned": int(o.get("Subtype/lineage not assigned", 0)),
            "HA only": int(tier.get("HA only", 0)),
            "HA + NA": int(tier.get("HA + NA", 0)),
            "Whole genome": int(tier.get("Whole genome", 0)),
            "No HA segment": int(o.get("No HA segment", 0)),
            "Failed QC / coverage": int(ha.get("qc", 0)),
            "Too short": int(ha.get("length", 0)),
            "Excluded": int(ha.get("excluded", 0)),
            "Clock outlier": int(ha.get("clock", 0) + ha.get("export", 0)),
            "HA in final build": int(o.get("In final build", 0)),
            "Genome in final build": int((g["genome_status"] == "in_build").sum()),
            "HA build yield": o.get("In final build", 0) / n if n else np.nan,
        })
    runs = pd.DataFrame(rows)

    prov = read_provenance(paths["run_logs"], args.run_id)
    if as_of is None:
        as_of = pd.Timestamp(dt.date.today())
    meta = {
        "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "as_of": as_of,
        "season_start": season_start_for(as_of),
        "prepared_by": prepared_by(),
        "provenance": prov,
        "paths": paths,
        "warnings": warnings,
        "data_through": lab["date"].max(),
    }
    return ReportData(cfg=cfg, meta=meta, seqs=seqs, lab=lab, samples=samples,
                      build_qc=build_qc, builds=builds, runs=runs, run_order=run_order,
                      run_dates={r: run_dates.loc[r, "min"] for r in run_order})


# ======================================================================
# Frequency data helpers
# ======================================================================


@app.function
def plot_window(d, df):
    """Rows of `df` inside the configured plotting window (lab specimens of
    the three built subtypes, dated, not in excluded runs)."""
    cfg = d.cfg
    out = df[df["subtype"].isin(SUBTYPE_ORDER) & df["date"].notna()]
    if cfg.get("plot_start_date"):
        out = out[out["date"] >= pd.Timestamp(cfg["plot_start_date"])]
    if cfg.get("plot_end_date"):
        out = out[out["date"] <= pd.Timestamp(cfg["plot_end_date"])]
    if cfg.get("exclude_runs_from_plots"):
        out = out[~out["sequencing_run"].isin(cfg["exclude_runs_from_plots"])]
    return out[out["sequencing_run"].notna()]


@app.function
def run_axis(df):
    """Sequencing runs ordered by date, labelled 'IV25Run3 · 2025-01-14'."""
    first = df.groupby("sequencing_run")["date"].min()
    order = sorted(first.index, key=lambda r: (first[r], natural_key(r)))
    labels = {r: f"{r} · {first[r]:%Y-%m-%d}" for r in order}
    return [labels[r] for r in order], labels


@app.function
def lineage_colors(values_by_count, top_n):
    """Top-N lineages get fixed palette slots (by abundance); the rest fold
    into 'Other'."""
    top = list(values_by_count[:top_n])
    domain = top + (["Other"] if len(values_by_count) > top_n else [])
    rng = PALETTE[:len(top)] + ([OTHER_COLOR] if len(values_by_count) > top_n else [])
    return top, domain, rng


@app.function
def stacked_run_chart(df, color_field, domain, colors, width, title_field, normalize=False,
                      tooltip_extra=None, legend_title=None):
    """Stacked bars per sequencing run (shared by subtype & lineage plots)."""
    order, labels = run_axis(df)
    data = (df.assign(run_label=df["sequencing_run"].map(labels))
            .groupby(["run_label", color_field], dropna=False).size().rename("n").reset_index())
    totals = data.groupby("run_label")["n"].transform("sum")
    data["share"] = data["n"] / totals
    data["total"] = totals
    data["order"] = data[color_field].map({v: i for i, v in enumerate(domain)})
    y = (alt.Y("share:Q", title="Share of sequences", axis=alt.Axis(format="%"),
               stack="normalize") if normalize else
         alt.Y("n:Q", title="Sequences", stack="zero"))
    base = alt.Chart(data).encode(
        x=alt.X("run_label:N", sort=order, title="Sequencing run · first specimen date",
                axis=alt.Axis(labelAngle=-50, labelLimit=160)),
    )
    bars = base.mark_bar().encode(
        y=y,
        color=alt.Color(f"{color_field}:N", title=legend_title or title_field,
                        scale=alt.Scale(domain=domain, range=colors),
                        legend=alt.Legend(orient="top", columns=min(len(domain), 8), offset=34)),
        order=alt.Order("order:Q"),
        tooltip=[alt.Tooltip("run_label:N", title="Run"),
                 alt.Tooltip(f"{color_field}:N", title=title_field),
                 alt.Tooltip("n:Q", title="Sequences"),
                 alt.Tooltip("share:Q", title="Share of run", format=".1%"),
                 alt.Tooltip("total:Q", title="Run total")],
    )
    layers = [bars]
    if not normalize:
        tot = data.drop_duplicates("run_label")
        layers.append(alt.Chart(tot).mark_text(dy=-6, fontSize=10, color=INK_2).encode(
            x=alt.X("run_label:N", sort=order), y=alt.Y("total:Q"), text="total:Q"))
    # Season markers (visual aid only): a dotted line just before the first
    # run on or after each Oct 1 / May 31.
    first = df.groupby("sequencing_run")["date"].min()
    runs_sorted = sorted(first.index, key=lambda r: (first[r], natural_key(r)))
    marks = {}
    for t, name in season_markers(first.min(), first.max()):
        nxt = next((r for r in runs_sorted if first[r] >= t), None)
        if nxt is not None and nxt != runs_sorted[0]:
            marks.setdefault(labels[nxt], []).append(name)
    if marks:
        mk = pd.DataFrame({"run_label": list(marks), "label": [" · ".join(v) for v in marks.values()]})
        # stagger labels of neighbouring markers onto two rows so they never overlap
        pos = {lab: i for i, lab in enumerate(order)}
        row, prev = [], None
        for lab in mk["run_label"]:
            r = 1 if (prev is not None and pos[lab] - pos[prev[0]] <= 3 and prev[1] == 0) else 0
            row.append(r)
            prev = (lab, r)
        mk["ypx"] = [-4 - 13 * (1 - r) for r in row]   # above the plot area
        xm = alt.X("run_label:N", sort=order, bandPosition=0)
        layers.append(alt.Chart(mk).mark_rule(color=MUTED, strokeWidth=1.2, strokeDash=[2, 3])
                      .encode(x=xm, tooltip=[alt.Tooltip("label:N", title="Season marker (visual aid)")]))
        layers.append(alt.Chart(mk).mark_text(align="left", baseline="bottom", dx=3, fontSize=10,
                                              color=MUTED, angle=0)
                      .encode(x=xm, y=alt.Y("ypx:Q", scale=None), text="label:N"))
    return alt.layer(*layers).properties(width=width, height=340)


# ======================================================================
# Report items. Each function returns one block (or a short list of blocks).
# Register new items in report_layout() at the end of this part.
# ======================================================================


# ---------- shared helpers for the items ----------------------------------


@app.function
def real_specimens(d):
    """Submitted specimens without controls."""
    return d.samples[~d.samples["is_control"]]


@app.function
def latest_run(d, season=None):
    """Most recent sequencing run (by first specimen date): overall when
    season is None, otherwise the last run whose first date is in `season`
    (a season start Timestamp). Returns (run, first_date) or (None, None)."""
    runs = [r for r in d.run_order if pd.notna(d.run_dates.get(r))]
    if season is not None:
        end = season_end(season)
        runs = [r for r in runs if season <= d.run_dates[r] < end]
    if not runs:
        return None, None
    return runs[-1], d.run_dates[runs[-1]]


@app.function
def run_columns(d, season):
    """(heading_latest, rows_latest, heading_season, rows_season) for the
    'Most recent run' and 'Season to date' columns of the season tables.
    Rows are JHH sequences in fludb (d.lab)."""
    is_current = season == d.meta["season_start"]
    run, rdate = latest_run(d, None if is_current else season)
    lab = d.lab
    end = season_end(season)
    in_season = lab[(lab["date"] >= season) & (lab["date"] < end)]
    if run is None:
        h_run, rows_run = "Most recent run (none)", lab.iloc[0:0]
    else:
        note = "" if season <= rdate < end else ", previous season"
        h_run = f"Most recent run: {run} ({rdate:%m/%d/%Y}{note})"
        rows_run = lab[lab["sequencing_run"] == run]
    upto = "to date" if is_current else "total"
    h_season = f"Season {upto}: {season_label(season)} (since {season:%b %d, %Y})"
    return h_run, rows_run, h_season, in_season


@app.function
def summary_table(d, season):
    """Influenza surveillance summary for one season (two count columns)."""
    h_run, g_run, h_season, g_season = run_columns(d, season)

    def col(g):
        n = len(g)
        a = int((g["type"] == "InfluenzaA").sum())
        h1 = int(((g["type"] == "InfluenzaA") & (g["subtype"] == "H1N1")).sum())
        h3 = int(((g["type"] == "InfluenzaA") & (g["subtype"] == "H3N2")).sum())
        b = int((g["type"] == "InfluenzaB").sum())
        vic = int(((g["type"] == "InfluenzaB") & (g["subtype"] == "Victoria")).sum())
        yam = int(((g["type"] == "InfluenzaB") & (g["subtype"] == "Yamagata")).sum())
        return [fmt_n(n), fmt_n_pct(a, n), fmt_n_pct(h1, a), fmt_n_pct(h3, a),
                fmt_n_pct(a - h1 - h3, a), fmt_n_pct(b, n), fmt_n_pct(vic, b), fmt_n_pct(yam, b)]

    labels = ["Total specimen sequences", "Influenza A", "A/H1N1pdm09", "A/H3N2",
              "Subtyping not possible", "Influenza B", "B/Victoria", "B/Yamagata"]
    styles = ["total", "total", "indent", "indent", "indent", "total", "indent", "indent"]
    table = pd.DataFrame({"": labels, h_run: col(g_run), h_season: col(g_season)})
    return table, styles


@app.function
def genetic_table(d, season):
    """HA subclade counts per subtype/lineage for one season (two count columns)."""
    h_run, g_run, h_season, g_season = run_columns(d, season)

    def assigned(g):
        return g[g["subclade"].notna() & (g["subclade"] != "")]

    g_run, g_season = assigned(g_run), assigned(g_season)
    rows, styles = [], []
    tree = [("Influenza A", "InfluenzaA", [("A/H1N1pdm09", "H1N1"), ("A/H3N2", "H3N2")]),
            ("Influenza B", "InfluenzaB", [("B/Victoria", "Victoria"), ("B/Yamagata", "Yamagata")])]
    for type_name, type_code, subs in tree:
        tr = g_run[g_run["type"] == type_code]
        ts = g_season[g_season["type"] == type_code]
        rows.append([type_name, "", fmt_n(len(tr)), fmt_n(len(ts))])
        styles.append("total")
        for sub_name, sub_code in subs:
            sr = tr[tr["subtype"] == sub_code]
            ss_ = ts[ts["subtype"] == sub_code]
            rows.append([sub_name, "", fmt_n(len(sr)), fmt_n(len(ss_))])
            styles.append("sub")
            both = pd.concat([sr, ss_])
            if both.empty:
                continue
            order = (ss_["subclade"].value_counts().reindex(both["subclade"].unique()).fillna(0)
                     .to_frame("s").assign(r=sr["subclade"].value_counts()).fillna(0)
                     .sort_values(["s", "r"], ascending=False).index)
            for sc in order:
                legacy = both.loc[both["subclade"] == sc, "legacy_clade"].mode() \
                    if "legacy_clade" in both else pd.Series(dtype=str)
                rows.append([sc, legacy.iloc[0] if len(legacy) else "",
                             fmt_n_pct((sr["subclade"] == sc).sum(), len(sr)),
                             fmt_n_pct((ss_["subclade"] == sc).sum(), len(ss_))])
                styles.append("indent2")
    table = pd.DataFrame(rows, columns=["Type / subtype / HA subclade", "HA clade (legacy)",
                                        h_run, h_season])
    return table, styles


@app.function
def run_label_map(d):
    return {r: f"{r} · {d.run_dates[r]:%Y-%m-%d}" if pd.notna(d.run_dates.get(r)) else r
            for r in d.run_order}


# ---------- Overview ----------------------------------------------------


@app.function
def item_at_a_glance(d):
    """Most recent run vs season to date: submitted vs in the final build."""
    s = real_specimens(d)
    ss = d.meta["season_start"]
    run, rdate = latest_run(d)

    def panel(g):
        n = len(g)
        built = g[g["outcome"] == "In final build"]
        mix = " · ".join(f"{SUBTYPE_LABEL[st]} {fmt_n((built['subtype'] == st).sum())}"
                         for st in SUBTYPE_ORDER)
        return [("Submitted", fmt_n(n), "specimens fed into the pipeline"),
                ("In the final build", fmt_n(len(built)),
                 (f"{fmt_pct(len(built), n)} of submitted · " + mix) if n else "no specimens yet")]

    season_rows = s[(s["date"] >= ss) & (s["date"] < season_end(ss))]
    run_rows = s[s["sequencing_run"] == run] if run else s.iloc[0:0]
    in_prev = run is not None and not (ss <= rdate < season_end(ss))
    panels = [
        (f"Most recent run · {run or 'none'}",
         (f"First specimen {rdate:%b %d, %Y}" + (" (previous season)" if in_prev else "")) if run else "",
         panel(run_rows)),
        (f"Season to date · {season_label(ss)}",
         f"Since {ss:%b %d, %Y}" + ("" if len(season_rows) else " · no runs yet this season"),
         panel(season_rows)),
    ]
    return Compare(panels, title="At a glance",
                   caption="'In the final build' = the specimen's HA is a tip of the final Nextstrain "
                           "HA tree. " + SUBMITTED_DEFINITION + " " + COMPLEMENT_STATEMENT)


@app.function
def item_notes(d):
    blocks = [Text(n.get("text", ""), kind=n.get("kind", "note"), pptx=True)
              for n in (d.cfg.get("notes") or [])]
    blocks += [Text(w, kind="warning") for w in d.meta["warnings"]]
    return blocks


# ---------- Surveillance summary ----------------------------------------


@app.function
def item_season_tables(d):
    """Summary + genetic characterization, one view per season (selector in
    HTML; current season in the deck plus one appendix slide per prior season)."""
    current = d.meta["season_start"]
    seasons = sorted({season_start_for(x) for x in d.lab["date"].dropna()} | {current}, reverse=True)
    seasons = [s for s in seasons if s <= current]
    options = []
    for s in seasons:
        summ, sst = summary_table(d, s)
        gen, gst = genetic_table(d, s)
        label = season_label(s) + (" (current season)" if s == current else "")
        options.append((label, [
            Table(f"surveillance_summary_{s.year}",
                  "Johns Hopkins Hospital System – Influenza Surveillance Summary", summ,
                  caption="Counts are JHH specimen sequences in fludb. Percentages: Influenza A and B of "
                          "the total; subtypes of Influenza A; lineages of Influenza B. 'Subtyping not "
                          "possible' = Influenza A without a complete H1N1 or H3N2 call from flusort.",
                  interactive=False, row_styles=sst),
            Table(f"genetic_characterization_{s.year}",
                  "Genetic Characterization", gen,
                  caption="HA subclade calls from Nextclade for specimens with an assigned subclade. "
                          "Subclade percentages are of the subtype/lineage row above.",
                  interactive=False, row_styles=gst),
        ]))
    return Choice("Season", options, appendix_title="Appendix – season {label}")


# ---------- Frequencies -------------------------------------------------


@app.function
def item_subtype_frequency(d):
    df = plot_window(d, d.lab)
    if df.empty:
        return Text("No dated JHH H1N1/H3N2/Victoria sequences in the plotting window.", kind="note")
    domain = [s for s in SUBTYPE_ORDER if s in set(df["subtype"])]
    colors = [SUBTYPE_COLORS[s] for s in domain]
    df = df.assign(subtype_label=df["subtype"].map(SUBTYPE_LABEL))
    dom_l = [SUBTYPE_LABEL[s] for s in domain]
    n_runs = df["sequencing_run"].nunique()
    mk = lambda norm: (lambda w: stacked_run_chart(df, "subtype_label", dom_l, colors, w,
                                                   "Subtype", normalize=norm))
    cap = (f"JHH sequences per sequencing run since {d.cfg['plot_start_date']}, by subtype/lineage "
           f"(flusort). Hover a bar for counts and shares. {SEASON_MARKER_NOTE}")
    return Tabs([
        ("Counts", [Chart("subtype_frequency", "Subtype frequency by sequencing run", mk(False),
                          caption=cap, min_width=24 * n_runs)]),
        ("Proportion", [Chart("subtype_proportion", "Subtype proportion by sequencing run", mk(True),
                              caption=cap, min_width=24 * n_runs, pptx=False)]),
    ])


@app.function
def item_lineage_frequency(d):
    df = plot_window(d, d.lab)
    tabs = []
    top_n = int(d.cfg["top_n_lineages"])
    for st in SUBTYPE_ORDER:
        g = df[(df["subtype"] == st) & df["subclade"].notna() & (df["subclade"] != "")]
        if g.empty:
            continue
        n_runs = g["sequencing_run"].nunique()
        views = []
        for field_, label in (("subclade", "Subclade"), ("legacy_clade", "Clade (legacy)")):
            if field_ not in g or g[field_].isna().all():
                continue
            counts = g[field_].value_counts()
            top, domain, colors = lineage_colors(list(counts.index), top_n)
            gg = g.assign(lineage=g[field_].where(g[field_].isin(top), "Other"))
            for norm in (False, True):
                key = f"{SUBTYPE_DIR[st]}_{slug(label)}_{'share' if norm else 'counts'}"
                title = f"{SUBTYPE_LABEL[st]} {label.lower()} {'proportion' if norm else 'frequency'}"
                chart = (lambda gg=gg, domain=domain, colors=colors, label=label, norm=norm:
                         (lambda w: stacked_run_chart(gg, "lineage", domain, colors, w, label,
                                                      normalize=norm, legend_title=label)))()
                folded = (f" Lineages outside the {top_n} most common are grouped as Other "
                          f"({len(counts) - top_n} lineages); see the tables for every lineage."
                          if len(counts) > top_n else "")
                views.append((f"{label} · {'share' if norm else 'counts'}",
                              [Chart(key, title, chart, caption=f"JHH {SUBTYPE_LABEL[st]} sequences per "
                                     f"run, coloured by HA {label.lower()} (Nextclade).{folded} "
                                     f"{SEASON_MARKER_NOTE}",
                                     min_width=24 * n_runs,
                                     pptx=(field_ == "subclade" and not norm))]))
        tabs.append((SUBTYPE_LABEL[st], [Tabs(views)]))
    if not tabs:
        return Text("No clade assignments in the plotting window.", kind="note")
    return Tabs(tabs)


# ---------- QC: run-level funnel ----------------------------------------

@app.function
def funnel_frames(d):
    """Long tables for the run-level funnel: specimens remaining at each step,
    and specimens removed at each step with the reason."""
    s = real_specimens(d).copy()
    ha = s["ha_status"]
    flags = {
        "Submitted": pd.Series(True, index=s.index),
        "Sequence data in fludb": s["has_sequence"],
        "Subtype/lineage assigned": s["has_sequence"] & s["build_subtype"].notna(),
        "HA in build input": ha.notna() & (ha != "absent"),
        "HA passed augur filter": ha.isin(["in_build", "clock", "export"]),
        "HA in final build": ha == "in_build",
        "Whole genome sequenced (8 segments)": s["genome_status"].notna() & (s["genome_status"] != "absent"),
        "Genome in final build": s["genome_status"] == "in_build",
    }

    def reason(r):
        if not r["has_sequence"]:
            return "No sequence data"
        if pd.isna(r["build_subtype"]):
            return "Subtype/lineage not assigned"
        h = r["ha_status"]
        if h == "absent" or pd.isna(h):
            return "No HA segment"
        return {"qc": "Failed Nextclade QC / coverage", "length": "HA too short",
                "excluded": "Manually excluded", "no_sequence": "Other augur filter",
                "filter_other": "Other augur filter", "clock": "Clock-filter outlier",
                "export": "Clock-filter outlier"}.get(h)

    s["reason"] = s.apply(reason, axis=1)
    keys = ["sequencing_run", "type_group", "subtype_group"]
    rem = []
    for i, (stage, f) in enumerate(flags.items()):
        g = s[f].groupby(keys, dropna=False).size().rename("n").reset_index()
        g["stage"], g["step"] = stage, i
        rem.append(g)
    remaining = pd.concat(rem, ignore_index=True)
    removed = (s[s["reason"].notna()].groupby(keys + ["reason"], dropna=False).size()
               .rename("n").reset_index())
    removed["o_reason"] = removed["reason"].map({r: i for i, r in enumerate(REMOVAL_REASONS)})
    return remaining.rename(columns={"sequencing_run": "run"}), \
        removed.rename(columns={"sequencing_run": "run"}), list(flags)


@app.function
def funnel_chart(d, field, width, interactive=True):
    remaining, removed, stages = funnel_frames(d)
    domain = TYPE_GROUPS if field == "type_group" else SUBTYPE_GROUPS
    colors = [GROUP_COLORS[g] for g in domain]
    runs = ["All runs"] + list(reversed(d.run_order))
    color = alt.Color(f"{field}:N", title="Type" if field == "type_group" else "Subtype / lineage",
                      scale=alt.Scale(domain=domain, range=colors),
                      legend=alt.Legend(orient="top", columns=3))
    if interactive:
        sel = alt.param(name=f"run_{field}", value="All runs",
                        bind=alt.binding_select(options=runs, name="Sequencing run  "))
        filt = (sel == "All runs") | (alt.datum.run == sel)
    else:
        sel, filt = None, None

    def base(df):
        c = alt.Chart(df)
        return c.transform_filter(filt) if filt is not None else c

    # concatenated views cannot fill a container: fixed width (scrolls on phones)
    w = width if isinstance(width, int) else 600
    remaining = remaining.assign(o=remaining[field].map({g: i for i, g in enumerate(domain)}))
    bars = base(remaining).transform_aggregate(
        n="sum(n)", groupby=["stage", "step", field, "o"]).mark_bar(height={"band": 0.72}).encode(
        y=alt.Y("stage:N", sort=stages, title=None, axis=alt.Axis(labelLimit=260)),
        x=alt.X("n:Q", title="Specimens", stack="zero"),
        color=color, order=alt.Order("o:Q"),
        tooltip=[alt.Tooltip("stage:N", title="Step"), alt.Tooltip(f"{field}:N", title="Group"),
                 alt.Tooltip("n:Q", title="Specimens", format=",")])
    labels = base(remaining).transform_aggregate(total="sum(n)", groupby=["stage", "step"]) \
        .transform_joinaggregate(sub="max(total)") \
        .transform_calculate(label="format(datum.total, ',') + '  (' + format(datum.total / datum.sub, '.0%') + ')'") \
        .mark_text(align="left", dx=6, color=INK_2).encode(
            y=alt.Y("stage:N", sort=stages), x="total:Q", text="label:N")
    top = alt.layer(bars, labels).properties(width=w, height=300,
                                             title="Specimens remaining at each step")
    removed = removed.assign(o=removed[field].map({g: i for i, g in enumerate(domain)}))
    out = base(removed).transform_aggregate(n="sum(n)", groupby=["reason", "o", field]).mark_bar(
        height={"band": 0.7}).encode(
        y=alt.Y("reason:N", sort=REMOVAL_REASONS, title=None, axis=alt.Axis(labelLimit=260)),
        x=alt.X("n:Q", title="Specimens removed", stack="zero", axis=alt.Axis(tickMinStep=1, format="d")),
        color=color, order=alt.Order("o:Q"),
        tooltip=[alt.Tooltip("reason:N", title="Removed because"),
                 alt.Tooltip(f"{field}:N", title="Group"), alt.Tooltip("n:Q", title="Specimens")]
    ).properties(width=w, height=230, title="Removed at each step, and why")
    ch = alt.vconcat(top, out, spacing=28).resolve_scale(color="shared")
    return ch.add_params(sel) if sel is not None else ch


@app.function
def item_run_funnel(d):
    cap = ("From submission to the final build for all runs or the run picked in the menu under the "
           "chart. Top: specimens still present after each step (label = total and share of "
           "submitted). Bottom: specimens removed at each step and the reason. HA steps follow the "
           "HA build; the last two rows follow the whole-genome build. " + SUBMITTED_DEFINITION)
    return Tabs([
        ("By subtype / lineage", [Chart(
            "qc_funnel_subtype", "Run-level funnel: in, removed, and in the final build",
            lambda w: funnel_chart(d, "subtype_group", w),
            static=lambda w: funnel_chart(d, "subtype_group", w, interactive=False),
            caption=cap)]),
        ("By type", [Chart(
            "qc_funnel_type", "Run-level funnel by type",
            lambda w: funnel_chart(d, "type_group", w),
            static=lambda w: funnel_chart(d, "type_group", w, interactive=False),
            caption=cap, pptx=False)]),
    ])


@app.function
def item_run_funnel_table(d):
    remaining, removed, stages = funnel_frames(d)
    allr = remaining.groupby(["stage", "step", "subtype_group"]).n.sum().reset_index().assign(run="All runs")
    df = pd.concat([allr, remaining], ignore_index=True)
    wide = df.pivot_table(index=["run", "step", "stage"], columns="subtype_group", values="n",
                          aggfunc="sum", fill_value=0).reset_index()
    groups = [g for g in SUBTYPE_GROUPS if g in wide.columns]
    wide["Total"] = wide[groups].sum(axis=1)
    sub = wide[wide["step"] == 0].set_index("run")["Total"]
    wide["% of submitted"] = [fmt_pct(t, sub.get(r, 0)) for r, t in zip(wide["run"], wide["Total"])]
    order = {r: i for i, r in enumerate(["All runs"] + list(reversed(d.run_order)))}
    wide = wide.sort_values(["run", "step"], key=lambda c: c.map(order) if c.name == "run" else c)
    lab = run_label_map(d)
    wide.insert(1, "Date", wide["run"].map(lambda r: f"{d.run_dates[r]:%Y-%m-%d}"
                                          if r in d.run_dates and pd.notna(d.run_dates[r]) else ""))
    wide = wide.drop(columns="step").rename(columns={"run": "Run", "stage": "Step"})
    for g in groups + ["Total"]:
        wide[g] = wide[g].astype(int)
    return Table("qc_funnel_table", "Run-level funnel table", wide.reset_index(drop=True),
                 caption="Specimens remaining after each step, per run and per subtype/lineage group "
                         "('All runs' first). Type the run name in the search box to see one run.",
                 page_length=16, pptx=False)


# ---------- QC: per sequencing run --------------------------------------


@app.function
def item_qc_by_run(d):
    s = real_specimens(d)
    if s.empty:
        return Text("No submitted specimens found.", kind="note")
    lab_map = run_label_map(d)
    order_l = [lab_map[r] for r in d.run_order]
    data = (s.assign(run_label=s["sequencing_run"].map(lab_map))
            .groupby(["run_label", "type_label", "outcome"]).size().rename("n").reset_index())
    data["o"] = data["outcome"].map({o: i for i, o in enumerate(OUTCOME_ORDER)})
    types = ["All types", "Influenza A", "Influenza B", "Untyped"]

    def chart(w, interactive=True):
        base = alt.Chart(data)
        if interactive:
            sel = alt.param(name="type_filter", value="All types",
                            bind=alt.binding_select(options=types, name="Type  "))
            base = base.add_params(sel).transform_filter(
                (sel == "All types") | (alt.datum.type_label == sel))
        return base.transform_aggregate(n="sum(n)", groupby=["run_label", "outcome", "o"]).mark_bar().encode(
            x=alt.X("run_label:N", sort=order_l, title="Sequencing run · date",
                    axis=alt.Axis(labelAngle=-50, labelLimit=160)),
            y=alt.Y("n:Q", title="Specimens"),
            color=alt.Color("outcome:N", title="Outcome",
                            scale=alt.Scale(domain=OUTCOME_ORDER, range=[OUTCOME_COLORS[o] for o in OUTCOME_ORDER]),
                            legend=alt.Legend(columns=3)),
            order=alt.Order("o:Q"),
            tooltip=[alt.Tooltip("run_label:N", title="Run"), alt.Tooltip("outcome:N", title="Outcome"),
                     alt.Tooltip("n:Q", title="Specimens")],
        ).properties(width=w, height=340)

    return Chart("qc_by_run", "Where each run's specimens ended up", lambda w: chart(w),
                 static=lambda w: chart(w, interactive=False), min_width=24 * len(order_l),
                 caption="Outcome of every submitted specimen, per sequencing run. 'In final build' "
                         "means the specimen's HA is a tip of the final HA tree. Use the Type menu "
                         "under the chart to show Influenza A or B only.")


@app.function
def item_qc_runs_table(d):
    table = d.runs.copy()
    table["HA build yield"] = table["HA build yield"].map(lambda v: f"{v:.0%}" if pd.notna(v) else "–")
    table = table.iloc[::-1].reset_index(drop=True)
    pptx_cols = ["Run", "Date", "Submitted", "A/H3N2", "A/H1N1pdm09", "B/Victoria", "HA only",
                 "HA + NA", "Whole genome", "Failed QC / coverage", "Clock outlier",
                 "HA in final build", "Genome in final build", "HA build yield"]
    return Table("qc_runs", "QC summary per sequencing run", table, pptx_columns=pptx_cols,
                 caption="One row per sequencing run (newest first); controls are not counted. "
                         "HA only = HA sequenced without NA; HA + NA = both, but not all 8 segments; "
                         "Whole genome = all 8 segments (segments uploaded to fludb). Removal columns "
                         "follow the HA build.", page_length=15, pptx_max_rows=14)


@app.function
def item_qc_run_matrix(d):
    s = real_specimens(d)
    lab_qc = d.build_qc[d.build_qc["sequence_ID"].isin(s["sequence_ID"])]
    if lab_qc.empty:
        return Text("No build inputs found.", kind="note")
    m = (lab_qc.assign(ok=lab_qc["status"] == "in_build")
         .groupby(["sequencing_run", "build_subtype", "segment"])
         .agg(input=("ok", "size"), final=("ok", "sum")).reset_index())
    m["retained"] = m["final"] / m["input"]
    m["Subtype"] = m["build_subtype"].map(lambda x: SUBTYPE_LABEL[DIR_SUBTYPE[x]])
    m["Build"] = m["segment"].map(lambda x: "Genome" if x == "genome" else x.upper())
    m["label"] = m["final"].astype(int).astype(str) + "/" + m["input"].astype(int).astype(str)
    build_order = [b.upper() for b in SEGMENTS] + ["Genome"]
    order = [r for r in d.run_order if r in set(m["sequencing_run"])]
    latest = order[-1]

    def matrix(w, interactive=True):
        base = alt.Chart(m)
        if interactive:
            pick = alt.param(name="run_pick", value=latest,
                             bind=alt.binding_select(options=order[::-1], name="Sequencing run  "))
            base = base.add_params(pick).transform_filter(alt.datum.sequencing_run == pick)
        else:
            base = base.transform_filter(alt.datum.sequencing_run == latest)
        enc = dict(x=alt.X("Build:N", sort=build_order, title=None, axis=alt.Axis(labelAngle=0, orient="top")),
                   y=alt.Y("Subtype:N", sort=[SUBTYPE_LABEL[x] for x in SUBTYPE_ORDER], title=None))
        rect = base.mark_rect(cornerRadius=3).encode(
            **enc, color=alt.Color("retained:Q", title="Share in final build",
                                   scale=alt.Scale(domain=[0, 1], range=[SEQ_RAMP[0], SEQ_RAMP[-2]]),
                                   legend=alt.Legend(format="%", orient="right", gradientLength=120)),
            tooltip=[alt.Tooltip("sequencing_run:N", title="Run"), "Subtype:N", "Build:N",
                     alt.Tooltip("input:Q", title="Build input"),
                     alt.Tooltip("final:Q", title="In final build"),
                     alt.Tooltip("retained:Q", title="Retained", format=".0%")])
        text = base.mark_text(fontSize=12).encode(
            **enc, text="label:N",
            color=alt.condition(alt.datum.retained > 0.55, alt.value("#ffffff"), alt.value(INK)))
        ch = alt.layer(rect, text).properties(width=w, height=150)
        return ch if interactive else ch.properties(title=latest)

    return Chart("qc_run_matrix", "Run detail: specimens retained in each build", lambda w: matrix(w),
                 static=lambda w: matrix(w, interactive=False),
                 caption="Retained / input for each segment and genome build, for the run picked in "
                         "the menu under the chart (latest run by default).")


# ---------- QC: per build -------------------------------------------------


@app.function
def build_stage_long(d):
    b = d.builds
    long = b.melt(id_vars=["Subtype", "Build", "Input"], value_vars=BUILD_STAGES,
                  var_name="Stage", value_name="n")
    long = long[long["n"] > 0].copy()
    long["o"] = long["Stage"].map({s: i for i, s in enumerate(BUILD_STAGES)})
    long["share"] = long["n"] / long["Input"]
    return long


@app.function
def item_build_heatmap(d):
    b = d.builds
    if b is None or b.empty:
        return Text("No build inputs found.", kind="note")
    build_order = [x.upper() for x in SEGMENTS] + ["Genome"]
    sub_order = [SUBTYPE_LABEL[s] for s in SUBTYPE_ORDER]

    def heat(w):
        data = b.copy()
        data["label"] = data["In final build"].astype(str) + "/" + data["Input"].astype(str)
        enc = dict(x=alt.X("Build:N", sort=build_order, title=None, axis=alt.Axis(labelAngle=0, orient="top")),
                   y=alt.Y("Subtype:N", sort=sub_order, title=None))
        rect = alt.Chart(data).mark_rect(cornerRadius=3).encode(
            **enc, color=alt.Color("Retained:Q", title="Share retained",
                                   scale=alt.Scale(domain=[0.5, 1], range=[SEQ_RAMP[0], SEQ_RAMP[-2]], clamp=True),
                                   legend=alt.Legend(format="%", orient="right", gradientLength=120)),
            tooltip=["Subtype:N", "Build:N", alt.Tooltip("Input:Q", title="Input"),
                     alt.Tooltip("JHH input:Q", title="JHH input"),
                     *[alt.Tooltip(f"{c}:Q") for c in BUILD_STAGES[1:]],
                     alt.Tooltip("In final build:Q"),
                     alt.Tooltip("Retained:Q", format=".1%")])
        text = alt.Chart(data).mark_text(fontSize=12).encode(
            **enc, text="label:N",
            color=alt.condition(alt.datum.Retained > 0.7, alt.value("#ffffff"), alt.value(INK)))
        return alt.layer(rect, text).properties(width=w, height=150)

    return Chart("qc_build_heatmap", "Retention in each Nextstrain build", heat,
                 caption="Sequences in the final Auspice tree / sequences entering the build "
                         "(JHH specimens and vaccine references). Hover for the stage-by-stage counts.")


@app.function
def item_build_losses(d):
    if d.builds is None or d.builds.empty:
        return None
    long = build_stage_long(d)
    data = long[long["Stage"] != "In final build"]
    build_order = [x.upper() for x in SEGMENTS] + ["Genome"]
    sub_order = [SUBTYPE_LABEL[s] for s in SUBTYPE_ORDER]

    def losses_facet(w):
        ww = w if isinstance(w, int) else None
        ch = alt.Chart(data).mark_bar().encode(
            y=alt.Y("Build:N", sort=build_order, title=None),
            x=alt.X("n:Q", title="Sequences removed"),
            color=alt.Color("Stage:N", title="Removed at",
                            scale=alt.Scale(domain=BUILD_STAGES[1:], range=BUILD_STAGE_COLORS[1:])),
            order=alt.Order("o:Q"),
            tooltip=["Subtype:N", "Build:N", "Stage:N", alt.Tooltip("n:Q", title="Sequences"),
                     alt.Tooltip("share:Q", title="Share of input", format=".1%")],
        ).properties(width=(ww // 3 - 40) if ww else 260, height=230)
        return ch.facet(column=alt.Column("Subtype:N", sort=sub_order, title=None))

    return Chart("qc_build_losses", "Where sequences were removed, by build", losses_facet,
                 caption="Sequences removed at each step: augur filter (Nextclade QC/coverage, minimum "
                         "length, manual exclusion) and augur refine's molecular-clock filter.")


@app.function
def item_build_table(d):
    if d.builds is None or d.builds.empty:
        return None
    table = d.builds.drop(columns=["subtype_dir", "segment"]).copy()
    table["Retained"] = table["Retained"].map(lambda v: f"{v:.1%}")
    return Table("qc_builds", "Stage counts per build", table,
                 caption="All 27 builds. 'Clock outlier' counts tips pruned by augur refine "
                         "(--clock-filter-iqd).", page_length=27)


# ---------- QC: Nextclade metrics -----------------------------------------


@app.function
def item_nextclade_status(d):
    q = d.build_qc[(d.build_qc["segment"] != "genome")].copy()
    if q.empty:
        return Text("No Nextclade QC results found.", kind="note")
    q["qc_overallStatus"] = q["qc_overallStatus"].fillna("no result")
    status_order = ["good", "mediocre", "bad", "no result"]
    status_colors = ["#0ca30c", "#fab219", "#d03b3b", "#c3c2b7"]
    q["Subtype"] = q["build_subtype"].map(lambda x: SUBTYPE_LABEL[DIR_SUBTYPE[x]])
    q["Segment"] = q["segment"].str.upper()
    counts = q.groupby(["Subtype", "Segment", "qc_overallStatus"]).size().rename("n").reset_index()
    counts["o"] = counts["qc_overallStatus"].map({s: i for i, s in enumerate(status_order)})
    seg_order = [s.upper() for s in SEGMENTS]
    sub_order = [SUBTYPE_LABEL[s] for s in SUBTYPE_ORDER]

    def status_chart(w):
        ww = w if isinstance(w, int) else None
        return alt.Chart(counts).mark_bar().encode(
            y=alt.Y("Segment:N", sort=seg_order, title=None),
            x=alt.X("n:Q", title="Sequences", stack="normalize", axis=alt.Axis(format="%")),
            color=alt.Color("qc_overallStatus:N", title="Nextclade QC status",
                            scale=alt.Scale(domain=status_order, range=status_colors)),
            order=alt.Order("o:Q"),
            tooltip=["Subtype:N", "Segment:N", alt.Tooltip("qc_overallStatus:N", title="Status"),
                     alt.Tooltip("n:Q", title="Sequences")],
        ).properties(width=(ww // 3 - 40) if ww else 260, height=220).facet(
            column=alt.Column("Subtype:N", sort=sub_order, title=None))

    return Chart("qc_nextclade_status", "Nextclade QC status by segment", status_chart,
                 caption="Share of build-input sequences with each Nextclade overall QC status. The "
                         f"filter keeps {' and '.join(d.cfg['qc_pass_status'])} sequences with coverage "
                         f"≥ {d.cfg['qc_min_coverage']:.0%}.")


@app.function
def item_ha_scatter(d):
    q = d.build_qc[d.build_qc["segment"] == "ha"].copy()
    if q.empty:
        return None
    q["Subtype"] = q["build_subtype"].map(lambda x: SUBTYPE_LABEL[DIR_SUBTYPE[x]])
    q["coverage"] = pd.to_numeric(q["coverage"], errors="coerce")
    q["qc_overallScore"] = pd.to_numeric(q["qc_overallScore"], errors="coerce")
    ha = q.dropna(subset=["coverage", "qc_overallScore"])
    ha = ha.assign(Result=ha["status"].map(lambda s: "Kept" if s in ("in_build", "clock", "export") else "Removed"))
    sub_order = [SUBTYPE_LABEL[s] for s in SUBTYPE_ORDER]

    def scatter(w):
        sel = alt.selection_point(fields=["Subtype"], bind="legend")
        pts = alt.Chart(ha).mark_point(filled=True, size=44, opacity=0.75, strokeWidth=1).encode(
            x=alt.X("coverage:Q", title="HA coverage", scale=alt.Scale(zero=False), axis=alt.Axis(format="%")),
            y=alt.Y("qc_overallScore:Q", title="Nextclade QC score (higher = worse)",
                    scale=alt.Scale(type="symlog"),
                    axis=alt.Axis(values=[0, 3, 10, 30, 100, 300, 1000, 3000])),
            color=alt.Color("Subtype:N", scale=alt.Scale(domain=sub_order, range=PALETTE[:3]),
                            legend=alt.Legend(title="Subtype (click to isolate)")),
            shape=alt.Shape("Result:N", scale=alt.Scale(domain=["Kept", "Removed"],
                                                        range=["circle", "cross"]),
                            legend=alt.Legend(title="augur filter")),
            opacity=alt.condition(sel, alt.value(0.8), alt.value(0.08)),
            tooltip=[alt.Tooltip("strain:N", title="Specimen"), alt.Tooltip("sequencing_run:N", title="Run"),
                     "Subtype:N", alt.Tooltip("coverage:Q", format=".1%"),
                     alt.Tooltip("qc_overallScore:Q", title="QC score", format=".1f"),
                     alt.Tooltip("qc_overallStatus:N", title="QC status"), "Result:N"],
        ).add_params(sel)
        rule = alt.Chart(pd.DataFrame({"x": [d.cfg["qc_min_coverage"]]})).mark_rule(
            color=MUTED, strokeWidth=1).encode(x="x:Q")
        return alt.layer(rule, pts).properties(width=w, height=320)

    return Chart("qc_ha_scatter", "HA coverage vs Nextclade QC score", scatter,
                 caption="Each point is one HA sequence; crosses were removed by augur filter. The "
                         "vertical line marks the minimum coverage. Hover to identify a specimen.",
                 pptx=False)


# ---------- QC: detailed tables -------------------------------------------


@app.function
def item_removed_table(d):
    dropped = d.build_qc[~d.build_qc["status"].isin(["in_build", "unknown"])].copy()
    dropped["Subtype"] = dropped["build_subtype"].map(lambda x: SUBTYPE_LABEL[DIR_SUBTYPE[x]])
    dropped["Build"] = dropped["segment"].map(lambda x: "Genome" if x == "genome" else x.upper())
    dropped["Stage"] = dropped["status"].map(lambda x: "augur refine (clock filter)"
                                             if x in ("clock", "export") else "augur filter")
    dropped["Coverage"] = pd.to_numeric(dropped["coverage"], errors="coerce").map(
        lambda v: f"{v:.1%}" if pd.notna(v) else "")
    dtab = dropped.rename(columns={"strain": "Strain", "sequencing_run": "Run", "date": "Date",
                                   "reason": "Reason", "qc_overallStatus": "QC status",
                                   "qc_overallScore": "QC score", "database_origin": "Origin"})
    dtab = dtab[["Strain", "Run", "Date", "Subtype", "Build", "Stage", "Reason", "QC status",
                 "QC score", "Coverage", "Origin"]].sort_values(["Run", "Strain", "Build"],
                                                               na_position="last")
    return Table("qc_removed", "Every sequence removed from a build, with the reason", dtab,
                 caption="One row per sequence × build that did not reach the final tree (segments "
                         "that were not sequenced are not listed). Search or filter by run, specimen "
                         "or reason.", pptx=False, page_length=15)


@app.function
def item_specimen_journey(d):
    s = real_specimens(d).copy()
    cols = {"sequence_ID": "Specimen", "sequencing_run": "Run", "date": "Date",
            "type_label": "Type", "subtype": "Subtype (flusort)", "segments_submitted": "Segments submitted",
            "segment_count": "Segments in fludb", "segment_tier": "HA/NA/genome",
            "outcome": "Outcome", "outcome_detail": "Detail", "subclade": "HA subclade"}
    s["date"] = s["date"].dt.strftime("%Y-%m-%d")
    for seg in BUILDS:
        s[f"{seg}_status"] = s[f"{seg}_status"].map(STATUS_SHORT).fillna("")
    jt = s.rename(columns=cols)
    jt = jt.rename(columns={f"{seg}_status": ("Genome" if seg == "genome" else seg.upper()) for seg in BUILDS})
    order = (["Specimen", "Run", "Date", "Type", "Subtype (flusort)", "Outcome"]
             + [("Genome" if x == "genome" else x.upper()) for x in BUILDS]
             + ["Detail", "HA subclade", "HA/NA/genome", "Segments submitted", "Segments in fludb"])
    jt = jt[[c for c in order if c in jt.columns]].sort_values(["Run", "Specimen"], ascending=[False, True])
    jt = jt.dropna(axis=1, how="all")
    short_code = {v: k for k, v in STATUS_SHORT.items()}
    jh = jt.copy()
    for c in [("Genome" if x == "genome" else x.upper()) for x in BUILDS]:
        jh[c] = jh[c].map(lambda v: f'<span class="st st-{short_code.get(v, "none")}">{htmlmod.escape(v)}</span>'
                          if v else "")
    legend = " · ".join(f"{STATUS_SHORT[k]} = {v.lower()}" for k, v in STATUS_LABEL.items())
    return Table("qc_specimens", "Specimen journey: status in each build", jt,
                 caption=f"One row per submitted specimen (controls excluded). Build columns: {legend}.",
                 pptx=False, page_length=15, html_df=jh)


@app.function
def item_database_glance(d):
    """Whole-database metrics (all runs), formerly the opening 'At a glance'."""
    real = real_specimens(d)
    in_ha = int((real["outcome"] == "In final build").sum())
    in_g = int((real["genome_status"] == "in_build").sum())
    first, last = d.lab["date"].min(), d.lab["date"].max()
    span = f"{first:%b %Y} – {last:%b %Y}" if pd.notna(first) else ""
    return KPIs([
        ("Specimens submitted (all runs)", fmt_n(len(real)),
         f"{fmt_n(real['sequencing_run'].nunique())} sequencing runs · {span}"),
        ("With sequence data", fmt_n(real["has_sequence"].sum()),
         fmt_pct(real["has_sequence"].sum(), len(real)) + " of submitted"),
        ("HA in the final builds", fmt_n(in_ha), fmt_pct(in_ha, len(real)) + " of submitted"),
        ("Genome in the final builds", fmt_n(in_g), fmt_pct(in_g, len(real)) + " of submitted"),
    ], title="Database at a glance (all runs)")


# ---------- Interactive tables --------------------------------------------


@app.function
def summary_by_subtype(df, by):
    def agg(g):
        n = len(g)
        has_c = g["legacy_clade"].notna().sum() if "legacy_clade" in g else 0
        has_s = (g["subclade"].notna() & (g["subclade"] != "")).sum()
        return pd.Series({
            "Total samples": n,
            "Unique clades (legacy)": g["legacy_clade"].nunique() if "legacy_clade" in g else 0,
            "Unique subclades": g["subclade"].nunique(),
            "Samples with clade": int(has_c),
            "Samples with subclade": int(has_s),
            "Clade coverage": f"{100 * has_c / n:.1f}%" if n else "–",
            "Subclade coverage": f"{100 * has_s / n:.1f}%" if n else "–",
        })
    return df.groupby(by).apply(agg, include_groups=False).reset_index()


@app.function
def lab_built_subtypes(d):
    lab = d.lab[d.lab["subtype"].isin(SUBTYPE_ORDER)].copy()
    lab["Subtype"] = lab["subtype"].map(SUBTYPE_LABEL)
    return lab


@app.function
def item_clade_summary(d):
    return Table("clade_summary_by_subtype", "Clade and subclade summary by subtype",
                 summary_by_subtype(lab_built_subtypes(d), ["Subtype"]),
                 caption="All JHH sequences of the three built subtypes.")


@app.function
def item_clade_summary_by_season(d):
    lab = lab_built_subtypes(d)
    t = summary_by_subtype(lab[lab["season"].notna()].rename(columns={"season": "Season"}),
                           ["Season", "Subtype"])
    return Table("clade_summary_by_season", "Seasonal clade and subclade summary", t,
                 caption="Seasons start October 1 (specimens from June–September count toward the "
                         "season that started the previous October).", page_length=15)


@app.function
def item_genome_completeness(d):
    g = d.lab.copy()

    def cat(r):
        if r["subtype"] in SUBTYPE_ORDER:
            return r["subtype"]
        if pd.isna(r["subtype"]) or r["subtype"] == "":
            return f"{r['type']}_unknown"
        return f"{r['type']}_{r['subtype']}"
    g["Subtype"] = g.apply(cat, axis=1)
    g["Status"] = np.where(g["subtype"].isin(SUBTYPE_ORDER), "Complete", "Incomplete")
    rows = []
    for (t, st, status), x in g.groupby(["type", "Subtype", "Status"]):
        rows.append({"Type": t, "Subtype": st, "Subtyping": status, "N": len(x),
                     "Mean segments": round(x["segment_count"].mean(), 1),
                     "Median segments": x["segment_count"].median(),
                     "Min": x["segment_count"].min(), "Max": x["segment_count"].max(),
                     "HA (n)": int((x["has_ha"] == "yes").sum()),
                     "HA (%)": f"{100 * (x['has_ha'] == 'yes').mean():.1f}%",
                     "HA+NA (n)": int((x["has_ha_and_na"] == "yes").sum()),
                     "HA+NA (%)": f"{100 * (x['has_ha_and_na'] == 'yes').mean():.1f}%",
                     "Complete genome (n)": int((x["segment_count"] == 8).sum()),
                     "Complete genome (%)": f"{100 * (x['segment_count'] == 8).mean():.1f}%"})
    t3 = pd.DataFrame(rows).sort_values(["Type", "N"], ascending=[True, False])
    return Table("genome_completeness", "Genome completeness by type and subtype", t3,
                 caption="Segments per JHH sequence in fludb. 'Incomplete' subtyping = flusort could not "
                         "assign both HA and NA (e.g. H1xx, xxN2) or the type only.")


@app.function
def item_all_sequences(d):
    full = d.seqs.copy()
    full["date"] = full["date"].dt.strftime("%Y-%m-%d")
    full = full.drop(columns=["type_label", "virus"], errors="ignore")
    return Table("all_sequences", "All sequences (fludb) with clade calls", full,
                 caption="Every sequence in fludb, including GISAID vaccine references (also written "
                         "to reports/report.tsv).", pptx=False, page_length=15)


# ---------- Methods -------------------------------------------------------


@app.function
def item_methods(d):
    cfg, meta = d.cfg, d.meta
    prov = meta["provenance"]
    lab = d.lab
    lines = [
        f"**Scope.** {COMPLEMENT_STATEMENT} {SUBMITTED_DEFINITION}",
        f"**Data.** JHH specimens sequenced by the Mostafa lab (fludb origin `{cfg['lab_origin']}`), "
        f"{fmt_n(len(lab))} sequences from {fmt_n(lab['sequencing_run'].nunique())} sequencing runs, "
        f"dated {lab['date'].min():%Y-%m-%d} to {lab['date'].max():%Y-%m-%d}. Frequency plots start "
        f"{cfg['plot_start_date']}" + (f" and end {cfg['plot_end_date']}" if cfg.get("plot_end_date") else "")
        + (f"; runs excluded from plots: {', '.join(cfg['exclude_runs_from_plots'])}"
           if cfg.get("exclude_runs_from_plots") else "") + ".",
        "**Typing.** Segments are typed and subtyped by BLAST against a reference HA/NA database "
        "(flusort). Only specimens assigned H1N1, H3N2 or B/Victoria enter the builds; partial calls "
        "(e.g. H1xx, xxN2) are 'Subtyping not possible'.",
        "**Clades.** HA clades and subclades are assigned with Nextclade using the latest datasets "
        "(downloaded at each run) and propagated to every segment of the same specimen.",
        f"**Quality control.** augur filter keeps sequences with Nextclade coverage ≥ "
        f"{cfg['qc_min_coverage']:.0%} and QC status {' or '.join(cfg['qc_pass_status'])}, above the "
        "segment's minimum length and not listed in config/exclude.tsv. augur refine then removes tips "
        "more than 4 interquartile distances from the molecular-clock expectation. Genome builds "
        "require all eight segments and ≥ 13,000 nt.",
        f"**Seasons.** A Northern Hemisphere season runs from October 1 to about May 31. The current "
        f"season starts October 1 of this year from October onward, otherwise October 1 of last year, "
        f"so June–September still show the season that just ended; summer specimens count toward that "
        f"season. Current season: {season_label(meta['season_start'])} (report date "
        f"{meta['as_of']:%Y-%m-%d}). {SEASON_MARKER_NOTE}",
    ]
    if prov.get("run_id"):
        tools = ", ".join(f"{k} {prov[k]}" for k in ("augur", "nextclade", "snakemake") if prov.get(k))
        lines.append(f"**Pipeline run** {prov.get('run_id')}"
                     + (f" ({prov.get('status')})" if prov.get("status") else "")
                     + (f", commit {prov['commit']}" if prov.get("commit") else "")
                     + (f"; {tools}" if tools else "") + ".")
    lines.append(f"**Prepared by** {meta['prepared_by']}. {ATTRIBUTION}. "
                 f"Source code: [{cfg['source_url']}]({cfg['source_url']})")
    return Text("\n\n".join(lines), pptx=True)


# ---------- Registry --------------------------------------------------------


@app.function
def report_layout():
    """The report, in order: (key, title, intro, [item functions]).
    Add a new plot or table by writing an item function and listing it here."""
    return [
        ("overview", "Overview", "",
         [item_at_a_glance, item_notes]),
        ("summary", "Surveillance summary",
         "Most recent run and season to date, in the format reported to public-health partners. "
         "Pick another season in the menu.",
         [item_season_tables]),
        ("subtypes", "Type and subtype frequency", "", [item_subtype_frequency]),
        ("lineages", "Clade and subclade frequency",
         "HA clade and subclade composition of each sequencing run, per subtype/lineage.",
         [item_lineage_frequency]),
        ("qc-funnel", "Quality control: run-level funnel",
         "How many specimens went in, where and why they were removed, and how many made it into "
         "the final Nextstrain build.",
         [item_run_funnel, item_run_funnel_table]),
        ("qc-runs", "Quality control: per sequencing run", "",
         [item_qc_by_run, item_qc_runs_table, item_qc_run_matrix]),
        ("qc-builds", "Quality control: per build", "",
         [item_build_heatmap, item_build_losses, item_build_table]),
        ("qc-nextclade", "Quality control: Nextclade metrics", "",
         [item_nextclade_status, item_ha_scatter]),
        ("qc-tables", "Quality control: detailed tables", "",
         [item_removed_table, item_specimen_journey, item_database_glance]),
        ("tables", "Interactive tables", "",
         [item_clade_summary, item_clade_summary_by_season, item_genome_completeness,
          item_all_sequences]),
        ("methods", "Methods and notes", "", [item_methods]),
    ]


@app.function
def build_sections(d, strict=False):
    out = []
    for key, title, intro, items in report_layout():
        blocks = []
        for fn in items:
            try:
                res = fn(d)
            except Exception as e:  # one broken item should not sink the report
                if strict:
                    raise
                print(f"WARNING: report item '{fn.__name__}' failed: {e}", file=sys.stderr)
                traceback.print_exc()
                res = Text(f"{fn.__name__} could not be built: {type(e).__name__}: {e}", kind="error")
            if res is None:
                continue
            blocks += res if isinstance(res, list) else [res]
        out.append(Section(key, title, blocks, intro))
    return out


# ======================================================================
# Rendering: shared helpers
# ======================================================================


@app.function
def iter_blocks(blocks):
    """Flatten Tabs and Choice (every view; used by the figure/table exporters)."""
    for b in blocks:
        if isinstance(b, Tabs):
            for _, inner in b.tabs:
                yield from iter_blocks(inner)
        elif isinstance(b, Choice):
            for _, inner in b.options:
                yield from iter_blocks(inner)
        else:
            yield b


@app.function
def chart_spec(block, width="container", static=False):
    fn = block.static if (static and block.static) else block.build
    return fn(width).to_dict()


@app.function
def png_bytes(block, width=1100, scale=2):
    import vl_convert as vlc
    spec = chart_spec(block, width=width, static=True)
    spec.setdefault("config", {})
    spec["config"]["background"] = "#ffffff"
    spec["config"]["font"] = "Helvetica, Arial, Liberation Sans, DejaVu Sans, sans-serif"
    # larger text for slides
    for k, v in {"axis": {"labelFontSize": 13, "titleFontSize": 14},
                 "legend": {"labelFontSize": 13, "titleFontSize": 14},
                 "header": {"labelFontSize": 14}}.items():
        spec["config"].setdefault(k, {}).update(v)
    return vlc.vegalite_to_png(spec, scale=scale)


@app.function
def md_inline(text):
    """Tiny markdown: **bold**, `code`, [text](url), paragraphs."""
    t = htmlmod.escape(text)
    t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
    t = re.sub(r"`(.+?)`", r"<code>\1</code>", t)
    t = re.sub(r"\[(.+?)\]\((https?://[^)]+)\)", r'<a href="\2">\1</a>', t)
    return "".join(f"<p>{p}</p>" for p in t.split("\n\n"))


@app.function
def md_plain(text):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    t = re.sub(r"`(.+?)`", r"\1", t)
    return re.sub(r"\[(.+?)\]\((.+?)\)", r"\2", t)


# ======================================================================
# Rendering: HTML
# ======================================================================


@app.function
def html_css():
    return """
:root{color-scheme:light;
 --page:#f9f9f7;--surface:#ffffff;--surface-2:#f3f2ee;--ink:#0b0b0b;--ink-2:#52514e;--muted:#898781;
 --line:rgba(11,11,11,.10);--accent:#2a78d6;--accent-ink:#1c5cab;--hero:#0d366b;--hero-ink:#ffffff;
 --warn-bg:#fff6e0;--warn-line:#eda100;--imp-bg:#fdecec;--imp-line:#d03b3b;--note-bg:#eaf2fd;--note-line:#2a78d6;
 --radius:12px;--shadow:0 1px 2px rgba(11,11,11,.06),0 4px 16px rgba(11,11,11,.04)}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){color-scheme:dark;
 --page:#0d0d0d;--surface:#1a1a19;--surface-2:#232321;--ink:#ffffff;--ink-2:#c3c2b7;--muted:#898781;
 --line:rgba(255,255,255,.10);--accent:#3987e5;--accent-ink:#86b6ef;--hero:#10223d;--hero-ink:#ffffff;
 --warn-bg:#2d2410;--warn-line:#c98500;--imp-bg:#2d1515;--imp-line:#e66767;--note-bg:#13233a;--note-line:#3987e5}}
:root[data-theme="dark"]{color-scheme:dark;
 --page:#0d0d0d;--surface:#1a1a19;--surface-2:#232321;--ink:#ffffff;--ink-2:#c3c2b7;--muted:#898781;
 --line:rgba(255,255,255,.10);--accent:#3987e5;--accent-ink:#86b6ef;--hero:#10223d;--hero-ink:#ffffff;
 --warn-bg:#2d2410;--warn-line:#c98500;--imp-bg:#2d1515;--imp-line:#e66767;--note-bg:#13233a;--note-line:#3987e5}
*{box-sizing:border-box}
html{scroll-behavior:smooth;scroll-padding-top:72px}
body{margin:0;background:var(--page);color:var(--ink);font:15px/1.55 system-ui,-apple-system,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;-webkit-text-size-adjust:100%}
a{color:var(--accent-ink)}
code{font:13px ui-monospace,SFMono-Regular,Menlo,monospace;background:var(--surface-2);padding:1px 5px;border-radius:5px}
.topbar{position:sticky;top:0;z-index:20;display:flex;align-items:center;gap:12px;padding:10px 16px;background:color-mix(in srgb,var(--page) 88%,transparent);backdrop-filter:blur(8px);border-bottom:1px solid var(--line)}
.topbar .name{font-weight:650;font-size:14px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;flex:1}
.btn{font:inherit;font-size:13px;color:var(--ink);background:var(--surface);border:1px solid var(--line);border-radius:999px;padding:6px 12px;cursor:pointer}
.btn:hover{background:var(--surface-2)}
.menu-btn{display:none}
.hero{background:var(--hero);color:var(--hero-ink);padding:40px 16px 32px}
.hero-inner{max-width:1240px;margin:0 auto}
.hero .eyebrow{font-size:12px;letter-spacing:.08em;text-transform:uppercase;opacity:.75;margin:0 0 8px}
.hero h1{font-size:clamp(24px,3.4vw,36px);line-height:1.15;margin:0 0 8px;font-weight:700;max-width:28ch}
.hero .sub{opacity:.85;margin:0 0 18px}
.chips{display:flex;flex-wrap:wrap;gap:8px}
.chip{font-size:12.5px;border:1px solid rgba(255,255,255,.25);border-radius:999px;padding:4px 10px;background:rgba(255,255,255,.06)}
.confidential{margin-top:18px;font-size:12.5px;opacity:.85;border-left:3px solid #fab219;padding-left:10px;max-width:90ch}
.layout{max-width:1240px;margin:0 auto;display:grid;grid-template-columns:220px minmax(0,1fr);gap:28px;padding:24px 16px 64px}
nav.toc{position:sticky;top:64px;align-self:start;max-height:calc(100vh - 80px);overflow:auto;font-size:13.5px}
nav.toc a{display:block;padding:6px 10px;border-radius:8px;color:var(--ink-2);text-decoration:none}
nav.toc a:hover{background:var(--surface-2);color:var(--ink)}
nav.toc a.active{background:var(--surface);color:var(--ink);box-shadow:inset 3px 0 0 var(--accent)}
main{min-width:0}
section.sec{margin:0 0 40px}
section.sec>h2{font-size:22px;margin:0 0 4px;letter-spacing:-.01em}
section.sec>.intro{color:var(--ink-2);margin:0 0 14px}
.card{background:var(--surface);border:1px solid var(--line);border-radius:var(--radius);box-shadow:var(--shadow);padding:18px;margin:14px 0}
.card h3{font-size:16px;margin:0 0 4px}
.card .cap{color:var(--ink-2);font-size:13.5px;margin:0 0 10px;max-width:95ch}
.scroll-x{overflow-x:auto;-webkit-overflow-scrolling:touch}
.chart{width:100%}
.chart .vega-bindings{display:flex;flex-wrap:wrap;gap:12px;margin-top:10px;font-size:13px;color:var(--ink-2)}
.chart .vega-bind select{font:inherit;color:var(--ink);background:var(--surface);border:1px solid var(--line);border-radius:8px;padding:5px 8px;margin-left:4px}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:12px;margin:14px 0}
.kpi{background:var(--surface);border:1px solid var(--line);border-radius:var(--radius);padding:14px 16px;box-shadow:var(--shadow)}
.kpi .l{font-size:12.5px;color:var(--ink-2)}
.kpi .v{font-size:28px;font-weight:700;letter-spacing:-.02em;line-height:1.2;margin:2px 0;overflow-wrap:anywhere}
.kpi .c{font-size:12.5px;color:var(--muted)}
.callout{border-left:4px solid var(--note-line);background:var(--note-bg);border-radius:8px;padding:10px 14px;margin:12px 0;font-size:14px}
.callout p{margin:0}
.callout.warning{border-color:var(--warn-line);background:var(--warn-bg)}
.callout.important,.callout.error{border-color:var(--imp-line);background:var(--imp-bg)}
.callout .k{font-weight:650;margin-right:6px}
.prose p{margin:0 0 10px;max-width:95ch;overflow-wrap:anywhere}
.tabs{margin:14px 0}
.tablist{display:flex;flex-wrap:wrap;gap:6px;border-bottom:1px solid var(--line);margin-bottom:4px}
.tablist button{font:inherit;font-size:13.5px;border:0;background:none;color:var(--ink-2);padding:8px 12px;border-bottom:2px solid transparent;cursor:pointer;margin-bottom:-1px}
.tablist button[aria-selected="true"]{color:var(--ink);border-bottom-color:var(--accent);font-weight:600}
.tabpanel[hidden]{display:none}
.tabs .tabs .tablist button{font-size:12.5px;padding:6px 10px}
table.static{border-collapse:collapse;width:100%;font-size:13.5px;font-variant-numeric:tabular-nums}
table.static th{text-align:left;font-weight:600;color:var(--ink-2);border-bottom:1px solid var(--line);padding:8px 10px;vertical-align:bottom}
table.static td{padding:7px 10px;border-bottom:1px solid var(--line)}
table.static td:not(:first-child),table.static th:not(:first-child){text-align:right}
table.static tr.group td{font-weight:650;background:var(--surface-2)}
.attribution{margin:0 0 14px;font-size:12px;letter-spacing:.02em;opacity:.7}
.scope{margin-top:14px;font-size:13px;opacity:.9;max-width:90ch}
.blocktitle{font-size:16px;margin:18px 0 0}
.compare{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:14px;margin:14px 0 6px}
.cmp-panel{background:var(--surface);border:1px solid var(--line);border-radius:var(--radius);box-shadow:var(--shadow);padding:16px 18px}
.cmp-panel .h{font-weight:700;font-size:17px}
.cmp-panel .s{color:var(--ink-2);font-size:13px;margin-bottom:6px}
.cmp-row{padding:10px 0;border-top:1px solid var(--line)}
.cmp-row .l{font-size:12.5px;color:var(--ink-2)}
.cmp-row .v{font-size:30px;font-weight:700;letter-spacing:-.02em;line-height:1.15}
.cmp-row .c{font-size:12.5px;color:var(--muted)}
.choice{margin:14px 0}
.choice-label{display:inline-flex;align-items:center;gap:8px;font-weight:600;font-size:14px}
.choice-label select{font:inherit;font-weight:400;color:var(--ink);background:var(--surface);border:1px solid var(--line);border-radius:8px;padding:6px 10px}
.choicepanel[hidden]{display:none}
table.static tr.total td{font-weight:700;border-top:1px solid color-mix(in srgb,var(--ink) 25%,transparent)}
table.static tr.sub td{font-weight:600}
table.static tr.sub td:first-child,table.static tr.indent td:first-child{padding-left:26px}
table.static tr.indent2 td:first-child{padding-left:46px}
table.static tr.indent2 td{color:var(--ink-2)}
table.static tr.indent td:first-child{padding-left:26px}
.itable-wrap{font-size:13px}
.itable-wrap table.dataTable{font-variant-numeric:tabular-nums}
.itable-wrap .dt-buttons .dt-button{font-size:12.5px}
.st{display:inline-block;min-width:34px;text-align:center;border-radius:6px;padding:0 6px;font-size:12px;font-weight:600;line-height:20px}
.st-in_build{background:#d9f2d9;color:#006300}.st-clock,.st-export{background:#fff0c7;color:#7a5200}
.st-qc,.st-length,.st-filter_other{background:#fde3d7;color:#8a3410}.st-excluded,.st-no_sequence{background:#f9dada;color:#8f1f1f}
.st-absent{color:var(--muted)}
html.dark .st-in_build{background:#123d12;color:#8fe08f}html.dark .st-clock,html.dark .st-export{background:#3d300d;color:#ffd77a}
html.dark .st-qc,html.dark .st-length,html.dark .st-filter_other{background:#43220f;color:#ffb48f}
html.dark .st-excluded,html.dark .st-no_sequence{background:#431515;color:#ff9c9c}
footer{max-width:1240px;margin:0 auto;padding:0 16px 40px;color:var(--muted);font-size:12.5px}
@media (max-width:900px){
 .layout{grid-template-columns:1fr;gap:0;padding-top:12px}
 .menu-btn{display:inline-block}
 nav.toc{position:fixed;top:52px;left:0;right:0;max-height:70vh;background:var(--surface);border-bottom:1px solid var(--line);padding:8px;z-index:19;display:none;box-shadow:var(--shadow)}
 nav.toc.open{display:block}
 .card{padding:14px}
 .hero{padding:28px 16px 24px}
}
@media print{.topbar,nav.toc{display:none}.layout{display:block}.card{break-inside:avoid;box-shadow:none}}
"""


@app.function
def html_js(dark_map):
    return """
const DARK_MAP = %s;
const specs = {};
document.querySelectorAll('script[type="application/json"][data-chart]').forEach(s => {
  specs[s.dataset.chart] = JSON.parse(s.textContent);
});
const views = {};
function isDark(){
  const t = document.documentElement.dataset.theme;
  if (t === 'dark') return true;
  if (t === 'light') return false;
  return window.matchMedia('(prefers-color-scheme: dark)').matches;
}
function themed(spec){
  if (!isDark()) return spec;
  let s = JSON.stringify(spec);
  for (const [a, b] of Object.entries(DARK_MAP)) s = s.split(a).join(b);
  return JSON.parse(s);
}
function visible(el){ return el.offsetParent !== null; }
async function embed(id, force){
  const el = document.getElementById(id);
  if (!el || (!force && el.dataset.done) || !visible(el)) return;
  el.dataset.done = '1';
  if (views[id]) { views[id].finalize(); }
  try {
    const res = await vegaEmbed(el, themed(specs[id]), {
      actions: {export: true, source: false, compiled: false, editor: false},
      renderer: 'svg', downloadFileName: id});
    views[id] = res;
  } catch (e) { el.textContent = 'Chart failed to render: ' + e; }
}
function embedVisible(force){ Object.keys(specs).forEach(id => embed(id, force)); }
function adjustTables(root){
  if (!window.jQuery || !jQuery.fn || !jQuery.fn.dataTable) return;
  jQuery(root).find('table.dataTable').each(function(){ try { jQuery(this).DataTable().columns.adjust(); } catch(e){} });
}
// Tabs
document.querySelectorAll('.tabs').forEach(tabs => {
  const list = tabs.querySelector(':scope > .tablist');
  const btns = list.querySelectorAll(':scope > button');
  const panels = tabs.querySelectorAll(':scope > .tabpanel');
  btns.forEach((b, i) => b.addEventListener('click', () => {
    btns.forEach((x, j) => x.setAttribute('aria-selected', i === j ? 'true' : 'false'));
    panels.forEach((p, j) => p.hidden = i !== j);
    embedVisible(false); adjustTables(panels[i]);
  }));
});
// Choice menus (e.g. season selector)
document.querySelectorAll('.choice').forEach(ch => {
  const sel = ch.querySelector(':scope > label > select');
  const panels = ch.querySelectorAll(':scope > .choicepanel');
  sel.addEventListener('change', () => {
    panels.forEach((p, j) => p.hidden = String(j) !== sel.value);
    embedVisible(false); adjustTables(ch);
  });
});
// Theme toggle
const themeBtn = document.getElementById('theme');
const modes = ['auto', 'light', 'dark'];
function setTheme(m){
  if (m === 'auto') delete document.documentElement.dataset.theme; else document.documentElement.dataset.theme = m;
  document.documentElement.classList.toggle('dark', isDark());
  themeBtn.textContent = 'Theme: ' + m;
  document.querySelectorAll('[data-done]').forEach(el => delete el.dataset.done);
  embedVisible(true);
}
themeBtn.addEventListener('click', () => {
  const cur = document.documentElement.dataset.theme || 'auto';
  setTheme(modes[(modes.indexOf(cur) + 1) %% 3]);
});
window.matchMedia('(prefers-color-scheme: dark)').addEventListener('change', () => {
  if (!document.documentElement.dataset.theme) setTheme('auto');
});
document.documentElement.classList.toggle('dark', isDark());
// Mobile menu + active section
const toc = document.querySelector('nav.toc');
document.getElementById('menu').addEventListener('click', () => toc.classList.toggle('open'));
toc.querySelectorAll('a').forEach(a => a.addEventListener('click', () => toc.classList.remove('open')));
const links = {}; toc.querySelectorAll('a').forEach(a => links[a.getAttribute('href').slice(1)] = a);
const obs = new IntersectionObserver(es => es.forEach(e => {
  if (e.isIntersecting) { Object.values(links).forEach(a => a.classList.remove('active'));
    const a = links[e.target.id]; if (a) a.classList.add('active'); }
}), {rootMargin: '-30%% 0px -60%% 0px'});
document.querySelectorAll('section.sec').forEach(s => obs.observe(s));
embedVisible(false);
""" % json.dumps(dark_map)


@app.function
def html_static_table(b):
    cols = list(b.df.columns)
    th = "".join(f"<th>{htmlmod.escape(str(c))}</th>" for c in cols)
    trs = []
    styles = b.row_styles or [None] * len(b.df)
    for i, row in enumerate(b.df.itertuples(index=False)):
        tds = "".join(f"<td>{htmlmod.escape('' if pd.isna(v) else str(v))}</td>" for v in row)
        trs.append(f'<tr class="{styles[i] or ""}">{tds}</tr>')
    return (f'<div class="scroll-x"><table class="static"><thead><tr>{th}</tr></thead>'
            f'<tbody>{"".join(trs)}</tbody></table></div>')


@app.function
def html_block(b, counter):
    counter[0] += 1
    if isinstance(b, Text):
        if b.kind == "text":
            return f'<div class="card prose">{md_inline(b.body)}</div>'
        label = {"warning": "Warning", "important": "Important", "error": "Error", "note": "Note"}.get(b.kind, "Note")
        return f'<div class="callout {b.kind}"><span class="k">{label}</span>{md_inline(b.body)[3:-4]}</div>'
    if isinstance(b, KPIs):
        tiles = "".join(
            f'<div class="kpi"><div class="l">{htmlmod.escape(l)}</div><div class="v">{htmlmod.escape(str(v))}</div>'
            f'<div class="c">{htmlmod.escape(c)}</div></div>' for l, v, c in b.items)
        return f'<h3 class="blocktitle">{htmlmod.escape(b.title)}</h3><div class="kpis">{tiles}</div>'
    if isinstance(b, Compare):
        panels = []
        for heading, sub, items in b.panels:
            rows = "".join(
                f'<div class="cmp-row"><div class="l">{htmlmod.escape(l)}</div>'
                f'<div class="v">{htmlmod.escape(str(v))}</div><div class="c">{htmlmod.escape(c)}</div></div>'
                for l, v, c in items)
            panels.append(f'<div class="cmp-panel"><div class="h">{htmlmod.escape(heading)}</div>'
                          f'<div class="s">{htmlmod.escape(sub)}</div>{rows}</div>')
        cap = f'<p class="cap">{htmlmod.escape(b.caption)}</p>' if b.caption else ""
        return f'<div class="compare">{"".join(panels)}</div>{cap}'
    if isinstance(b, Chart):
        cid = f"chart_{b.key}"
        spec = json.dumps(chart_spec(b), default=str).replace("</", "<\\/")
        style = f' style="min-width:{b.min_width}px"' if b.min_width else ""
        return (f'<div class="card"><h3>{htmlmod.escape(b.title)}</h3>'
                f'<p class="cap">{htmlmod.escape(b.caption)}</p>'
                f'<div class="scroll-x"><div class="chart" id="{cid}"{style}></div></div>'
                f'<script type="application/json" data-chart="{cid}">{spec}</script></div>')
    if isinstance(b, Table):
        head = (f'<div class="card"><h3>{htmlmod.escape(b.title)}</h3>'
                + (f'<p class="cap">{htmlmod.escape(b.caption)}</p>' if b.caption else ""))
        if not b.interactive:
            return head + html_static_table(b) + "</div>"
        from itables import to_html_datatable
        df = (b.html_df if b.html_df is not None else b.df).copy()
        for c in df.columns:
            if pd.api.types.is_datetime64_any_dtype(df[c]):
                df[c] = df[c].dt.strftime("%Y-%m-%d")
        df = df.astype(object).where(df.notna(), "")
        tbl = to_html_datatable(
            df, connected=False, table_id=f"tbl_{b.key}", showIndex=False, maxBytes=0,
            classes="display compact nowrap", scrollX=True, pageLength=b.page_length,
            lengthMenu=[10, 15, 25, 50, 100, -1],
            layout={"topStart": "buttons", "topEnd": "search", "bottomStart": "info",
                    "bottomEnd": "paging"},
            buttons=["pageLength", "copyHtml5", "csvHtml5", "colvis"],
            columnControl=["order", ["searchList", "spacer", "orderAsc", "orderDesc", "orderClear"]],
            ordering={"indicators": False, "handler": False},
            style="width:100%;margin:0", allow_html=b.html_df is not None,
            **({"fixedColumns": {"start": 1}} if len(df.columns) > 8 else {}),
        )
        return head + f'<div class="itable-wrap">{tbl}</div></div>'
    if isinstance(b, Tabs):
        btns, panels = [], []
        for i, (label, inner) in enumerate(b.tabs):
            sel = "true" if i == 0 else "false"
            btns.append(f'<button role="tab" aria-selected="{sel}">{htmlmod.escape(label)}</button>')
            body = "".join(html_block(x, counter) for x in inner)
            panels.append(f'<div class="tabpanel" role="tabpanel"{"" if i == 0 else " hidden"}>{body}</div>')
        return f'<div class="tabs"><div class="tablist" role="tablist">{"".join(btns)}</div>{"".join(panels)}</div>'
    if isinstance(b, Choice):
        opts = "".join(f'<option value="{i}">{htmlmod.escape(label)}</option>'
                       for i, (label, _) in enumerate(b.options))
        panels = "".join(
            f'<div class="choicepanel"{"" if i == 0 else " hidden"}>{"".join(html_block(x, counter) for x in inner)}</div>'
            for i, (_, inner) in enumerate(b.options))
        return (f'<div class="choice"><label class="choice-label">{htmlmod.escape(b.label)} '
                f'<select>{opts}</select></label>{panels}</div>')
    return ""


@app.function
def render_html(d, sections):
    import vl_convert as vlc
    from itables.javascript import generate_init_offline_itables_html
    from itables.utils import find_package_file

    cfg, meta = d.cfg, d.meta
    prov = meta["provenance"]
    counter = [0]
    body = []
    toc = []
    for s in sections:
        toc.append(f'<a href="#{s.key}">{htmlmod.escape(s.title)}</a>')
        inner = "".join(html_block(b, counter) for b in s.blocks)
        intro = f'<p class="intro">{htmlmod.escape(s.intro)}</p>' if s.intro else ""
        body.append(f'<section class="sec" id="{s.key}"><h2>{htmlmod.escape(s.title)}</h2>{intro}{inner}</section>')
    chips = [f"Prepared by {meta['prepared_by']}",
             f"Generated {meta['generated']}",
             f"Data through {meta['data_through']:%Y-%m-%d}" if pd.notna(meta["data_through"]) else None,
             f"Current season {season_label(meta['season_start'])}",
             f"Pipeline run {prov['run_id']}" if prov.get("run_id") else None,
             f"Commit {prov['commit']}" if prov.get("commit") else None]
    chips_html = "".join(f'<span class="chip">{htmlmod.escape(c)}</span>' for c in chips if c)
    dark_map = {a: b for a, b in zip(PALETTE, PALETTE_DARK)}
    dark_map.update({INK: "#ffffff", INK_2: "#c3c2b7", GRID: "#2c2c2a", SURFACE: "#1a1a19",
                     SEQ_RAMP[0]: "#1f2a38"})
    itables_init = generate_init_offline_itables_html(find_package_file("html/dt_bundle.js"))
    vega_bundle = vlc.javascript_bundle()
    title = htmlmod.escape(cfg["title"])
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Influenza Surveillance Report</title>
<meta name="description" content="{title}">
<meta name="author" content="{htmlmod.escape(ATTRIBUTION)}">
<style>{html_css()}</style>
{itables_init}
<script type="module">{vega_bundle}
window.dispatchEvent(new Event('vega-ready'));</script>
</head><body>
<div class="topbar"><button class="btn menu-btn" id="menu" aria-label="Sections">☰</button>
<div class="name">{title}</div><button class="btn" id="theme">Theme: auto</button></div>
<header class="hero"><div class="hero-inner"><p class="attribution">{htmlmod.escape(ATTRIBUTION)}</p>
<p class="eyebrow">{htmlmod.escape(cfg.get('author') or '')} · Nextstrain surveillance</p>
<h1>{title}</h1><p class="sub">{htmlmod.escape(cfg.get('subtitle') or '')}</p>
<div class="chips">{chips_html}</div>
<div class="scope">{htmlmod.escape(COMPLEMENT_STATEMENT + " " + SUBMITTED_DEFINITION)}</div>
<div class="confidential">{htmlmod.escape(cfg.get('confidentiality') or '')}</div></div></header>
<div class="layout"><nav class="toc" aria-label="Sections">{''.join(toc)}</nav>
<main>{''.join(body)}</main></div>
<footer>{htmlmod.escape(COMPLEMENT_STATEMENT)} {htmlmod.escape(SUBMITTED_DEFINITION)}<br>
Prepared by {htmlmod.escape(meta['prepared_by'])} · built by scripts/report.py · {htmlmod.escape(meta['generated'])} · {htmlmod.escape(ATTRIBUTION)}</footer>
<script type="module">
function start(){{ {html_js(dark_map)} }}
if (window.vegaEmbed) start(); else window.addEventListener('vega-ready', start, {{once: true}});
</script>
</body></html>"""


# ======================================================================
# Rendering: PowerPoint
# ======================================================================


@app.function
def render_pptx(d, sections, path, figure_cache=None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt
    from PIL import Image

    figure_cache = figure_cache or {}
    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))
    W, H = Inches(13.333), Inches(7.5)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    cfg, meta = d.cfg, d.meta
    footer_text = f"CONFIDENTIAL · prepared by {meta['prepared_by']} · generated {meta['generated']}"
    margin = Inches(0.55)
    content_top = Inches(1.35)
    content_h = H - content_top - Inches(0.95)
    per_slide = int(cfg.get("pptx_table_rows", 14))

    def textbox(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.02)
        paras = text if isinstance(text, list) else [text]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = rgb(color)
            r.font.name = "Calibri"
        return tb

    def base_slide(title, kicker=None):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb("#0d366b"); bar.line.fill.background()
        if kicker:
            textbox(s, margin, Inches(0.3), W - 2 * margin, Inches(0.3), kicker.upper(), 10.5,
                    color="#2a78d6", bold=True)
        textbox(s, margin, Inches(0.55), W - 2 * margin, Inches(0.7), title, 24, bold=True)
        textbox(s, margin, H - Inches(0.45), W - 2 * margin - Inches(1), Inches(0.3), footer_text, 9, color=MUTED)
        textbox(s, W - margin - Inches(1), H - Inches(0.45), Inches(1), Inches(0.3),
                str(len(prs.slides)), 9, color=MUTED, align=PP_ALIGN.RIGHT)
        return s

    def table_shape(s, df, styles, x, y, w, font, row_h, numeric_right=True, head_h=None):
        ncol = len(df.columns)
        head_h = head_h or max(row_h, Inches(0.42))
        shape = s.shapes.add_table(len(df) + 1, ncol, x, y, w, head_h + row_h * len(df))
        t = shape.table
        t.rows[0].height = int(head_h)
        for i in range(1, len(df) + 1):
            t.rows[i].height = int(row_h)
        if styles is not None and ncol >= 3:  # hierarchy tables: wide label column
            t.columns[0].width = int(w * (0.34 if ncol == 3 else 0.28))
            rest = int((w - t.columns[0].width) // (ncol - 1))
            for j in range(1, ncol):
                t.columns[j].width = rest
        for j, c in enumerate(df.columns):
            cell = t.cell(0, j)
            cell.text = str(c) if str(c).strip() else " "
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb("#0d366b")
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(font); para.runs[0].font.bold = True
            para.runs[0].font.color.rgb = rgb("#ffffff")
            if j > 0 and styles is not None:
                para.alignment = PP_ALIGN.RIGHT
        indent = {"indent": "    ", "sub": "    ", "indent2": "        "}
        for i, row in enumerate(df.itertuples(index=False)):
            style = styles[i] if styles is not None and i < len(styles) else None
            for j, v in enumerate(row):
                cell = t.cell(i + 1, j)
                txt = "" if (v is None or (isinstance(v, float) and np.isnan(v))) else str(v)
                if j == 0 and style in indent:
                    txt = indent[style] + txt
                cell.text = txt if txt else " "     # an empty cell would fall back to 18 pt
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb("#eef2f8" if style == "total" else
                                               ("#ffffff" if i % 2 == 0 else "#f7f7f5"))
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(font)
                if para.runs:
                    para.runs[0].font.size = Pt(font)
                    para.runs[0].font.bold = style in ("total", "sub")
                    para.runs[0].font.color.rgb = rgb(INK if style != "indent2" else INK_2)
                if j > 0 and (styles is not None or numeric_right) and (styles is not None):
                    para.alignment = PP_ALIGN.RIGHT
                cell.margin_top = cell.margin_bottom = Inches(0.01)
        return head_h + row_h * len(df)

    def table_slides(b, kicker):
        df = b.df[b.pptx_columns] if b.pptx_columns else b.df
        styles = b.row_styles
        caption = b.caption
        if b.pptx_max_rows and len(df) > b.pptx_max_rows:
            caption = (f"First {b.pptx_max_rows} of {len(df)} rows; the full table is in the HTML "
                       f"report and report.xlsx. " + caption)
            df = df.head(b.pptx_max_rows)
            styles = styles[: b.pptx_max_rows] if styles else None
        per = per_slide if not styles else 22
        chunks = [(df.iloc[i:i + per], (styles[i:i + per] if styles else None))
                  for i in range(0, max(len(df), 1), per)]
        for ci, (chunk, st) in enumerate(chunks):
            title = b.title + (f" ({ci + 1}/{len(chunks)})" if len(chunks) > 1 else "")
            s = base_slide(title, kicker)
            ncol = len(chunk.columns)
            font = 12 if ncol <= 5 else 10.5 if ncol <= 9 else 9
            if styles:
                font, row_h = (12, Inches(0.36)) if len(chunk) <= 12 else (10, Inches(0.235))
            else:
                row_h = Inches(0.36 if font >= 12 else 0.32)
            width = W - 2 * margin if not styles or ncol > 4 else Inches(9.5)
            h = table_shape(s, chunk, st, margin, content_top, width, font, row_h)
            if caption and ci == len(chunks) - 1 and content_top + h + Inches(0.15) < H - Inches(1.1):
                textbox(s, margin, content_top + h + Inches(0.15), W - 2 * margin, Inches(0.5),
                        caption, 11, color=INK_2)

    def chart_slide(b, kicker):
        png = figure_cache.get(b.key) or png_bytes(b)
        s = base_slide(b.title, kicker)
        im = Image.open(io.BytesIO(png))
        cap_h = Inches(0.6) if b.caption else 0
        max_w, max_h = W - 2 * margin, content_h - cap_h
        ratio = min(max_w / im.width, max_h / im.height)
        pw, ph = int(im.width * ratio), int(im.height * ratio)
        s.shapes.add_picture(io.BytesIO(png), int(margin + (max_w - pw) / 2), content_top, pw, ph)
        if b.caption:
            textbox(s, margin, content_top + ph + Inches(0.1), W - 2 * margin, cap_h, b.caption, 11, color=INK_2)

    def compare_slide(b, kicker):
        s = base_slide(b.title, kicker)
        n = len(b.panels)
        gap = Inches(0.35)
        pw = (W - 2 * margin - gap * (n - 1)) / n
        ph = Inches(4.3)
        for i, (heading, sub, items) in enumerate(b.panels):
            x = margin + i * (pw + gap)
            box = s.shapes.add_shape(1, x, content_top, pw, ph)
            box.fill.solid(); box.fill.fore_color.rgb = rgb("#f3f2ee"); box.line.fill.background()
            textbox(s, x + Inches(0.3), content_top + Inches(0.2), pw - Inches(0.6), Inches(0.45), heading, 20, bold=True)
            textbox(s, x + Inches(0.3), content_top + Inches(0.68), pw - Inches(0.6), Inches(0.35), sub, 12, color=INK_2)
            for k, (label, value, cap) in enumerate(items):
                y = content_top + Inches(1.2) + k * Inches(1.5)
                textbox(s, x + Inches(0.3), y, pw - Inches(0.6), Inches(0.3), label, 13, color=INK_2)
                textbox(s, x + Inches(0.3), y + Inches(0.28), pw - Inches(0.6), Inches(0.7), str(value), 36, bold=True)
                textbox(s, x + Inches(0.3), y + Inches(0.95), pw - Inches(0.6), Inches(0.45), cap, 11, color=MUTED)
        if b.caption:
            textbox(s, margin, content_top + ph + Inches(0.15), W - 2 * margin, Inches(0.6), b.caption, 11, color=INK_2)

    def kpi_slide(b, kicker):
        s = base_slide(b.title, kicker)
        n = len(b.items)
        cols = 3 if n > 4 else n
        gap = Inches(0.25)
        tw = (W - 2 * margin - gap * (cols - 1)) / cols
        th = Inches(1.7)
        for i, (label, value, cap) in enumerate(b.items):
            x = margin + (i % cols) * (tw + gap)
            y = content_top + Inches(0.2) + (i // cols) * (th + gap)
            box = s.shapes.add_shape(1, x, y, tw, th)
            box.fill.solid(); box.fill.fore_color.rgb = rgb("#f3f2ee"); box.line.fill.background()
            textbox(s, x + Inches(0.2), y + Inches(0.15), tw - Inches(0.4), Inches(0.3), label, 12, color=INK_2)
            textbox(s, x + Inches(0.2), y + Inches(0.45), tw - Inches(0.4), Inches(0.7), str(value),
                    30 if len(str(value)) < 12 else 22, bold=True)
            textbox(s, x + Inches(0.2), y + Inches(1.15), tw - Inches(0.4), Inches(0.45), cap, 11, color=MUTED)

    def text_slide(b, kicker, title=None):
        s = base_slide(title or ("Notes" if b.kind != "text" else kicker), kicker)
        paras = [md_plain(p) for p in b.body.split("\n\n")]
        if b.kind != "text":
            paras = [f"{b.kind.capitalize()}: {p}" for p in paras]
        textbox(s, margin, content_top, W - 2 * margin, content_h, paras, 12 if len(b.body) > 1100 else 14)

    def season_appendix_slide(title, blocks):
        """Summary table (left) and genetic characterization (right) on one
        slide; a long genetic table continues on extra slides in two columns."""
        tables = [b for b in blocks if isinstance(b, Table)]
        if not tables:
            return
        summ, gen = tables[0], (tables[1] if len(tables) > 1 else None)
        cap = 22                                  # genetic rows per column
        row_h, font = Inches(0.19), 8
        s = base_slide(title, "Appendix")
        left_w, right_x = Inches(4.9), margin + Inches(5.2)
        right_w = W - margin - right_x
        textbox(s, margin, content_top - Inches(0.05), left_w, Inches(0.3), summ.title, 11, bold=True)
        table_shape(s, summ.df, summ.row_styles, margin, content_top + Inches(0.3), left_w, 10.5, Inches(0.3))
        if gen is None:
            return
        df, st = gen.df, gen.row_styles or [None] * len(gen.df)
        chunks = [(df.iloc[i:i + cap], st[i:i + cap]) for i in range(0, len(df), cap)]
        textbox(s, right_x, content_top - Inches(0.05), right_w, Inches(0.3),
                gen.title + (f" (1/{len(chunks)})" if len(chunks) > 1 else ""), 11, bold=True)
        table_shape(s, chunks[0][0], chunks[0][1], right_x, content_top + Inches(0.3), right_w, font, row_h)
        rest = chunks[1:]
        for k in range(0, len(rest), 2):
            s = base_slide(title + " (continued)", "Appendix")
            colw = int((W - 2 * margin - Inches(0.3)) / 2)
            for j, (cdf, cst) in enumerate(rest[k:k + 2]):
                x = margin + j * (colw + Inches(0.3))
                textbox(s, x, content_top - Inches(0.05), colw, Inches(0.3),
                        f"{gen.title} ({k + j + 2}/{len(chunks)})", 11, bold=True)
                table_shape(s, cdf, cst, x, content_top + Inches(0.3), colw, font, row_h)

    # Title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = rgb("#0d366b"); bg.line.fill.background()
    textbox(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(0.35), ATTRIBUTION, 11, color="#9ec5f4")
    textbox(s, Inches(0.9), Inches(1.6), Inches(11.5), Inches(0.4), "NEXTSTRAIN SURVEILLANCE REPORT", 13,
            bold=True, color="#86b6ef")
    textbox(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.8), cfg["title"], 36, bold=True, color="#ffffff")
    textbox(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(0.5), cfg.get("subtitle") or "", 16, color="#cde2fb")
    prov = meta["provenance"]
    facts = [f"Prepared by {meta['prepared_by']}", f"Generated {meta['generated']}",
             f"Data through {meta['data_through']:%Y-%m-%d}" if pd.notna(meta["data_through"]) else "",
             f"Pipeline run {prov['run_id']}" if prov.get("run_id") else ""]
    textbox(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(0.4), "   ·   ".join(f for f in facts if f), 13,
            color="#cde2fb")
    textbox(s, Inches(0.9), Inches(5.4), Inches(11.5), Inches(0.6),
            COMPLEMENT_STATEMENT + " " + SUBMITTED_DEFINITION, 12, color="#cde2fb")
    textbox(s, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.8), cfg.get("confidentiality") or "", 10.5,
            color="#9ec5f4")

    appendix = []

    def emit(b, kicker):
        if isinstance(b, Tabs):
            for _, inner in b.tabs:
                for x in inner:
                    emit(x, kicker)
        elif isinstance(b, Choice):
            for x in b.options[0][1]:
                emit(x, kicker)
            for label, inner in b.options[1:]:
                appendix.append((b.appendix_title.format(label=label), inner))
        elif isinstance(b, Compare) and b.pptx:
            compare_slide(b, kicker)
        elif isinstance(b, KPIs) and b.pptx:
            kpi_slide(b, kicker)
        elif isinstance(b, Text) and b.pptx:
            text_slide(b, kicker)
        elif isinstance(b, Chart) and b.pptx:
            chart_slide(b, kicker)
        elif isinstance(b, Table) and b.pptx:
            table_slides(b, kicker)

    for sec in sections:
        for b in sec.blocks:
            emit(b, sec.title)
    for title, blocks in appendix:
        season_appendix_slide(title, blocks)
    prs.save(path)


# ======================================================================
# Writing everything
# ======================================================================


@app.function
def verify_attribution(html_path=None, pptx_path=None):
    """Fail the build if the attribution is missing from either output."""
    missing = []
    if html_path is not None and ATTRIBUTION not in Path(html_path).read_text(encoding="utf-8"):
        missing.append(str(html_path))
    if pptx_path is not None:
        from pptx import Presentation
        texts = [sh.text_frame.text for sl in Presentation(pptx_path).slides
                 for sh in sl.shapes if sh.has_text_frame]
        if not any(ATTRIBUTION in t for t in texts):
            missing.append(str(pptx_path))
    if missing:
        raise RuntimeError(f"Attribution '{ATTRIBUTION}' is missing from: {', '.join(missing)}. "
                           "It must not be removed (see the authorship note in scripts/report.py).")


@app.function
def write_outputs(d, sections, outdir, formats):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    written = []
    blocks = [b for s in sections for b in iter_blocks(s.blocks)]
    figures = {}
    if "png" in formats or "pptx" in formats:
        for b in blocks:
            if isinstance(b, Chart) and (b.pptx or "png" in formats):
                try:
                    figures[b.key] = png_bytes(b)
                except Exception as e:
                    print(f"WARNING: could not render {b.key} to PNG: {e}", file=sys.stderr)
    if "png" in formats:
        (outdir / "figures").mkdir(exist_ok=True)
        for k, png in figures.items():
            (outdir / "figures" / f"{k}.png").write_bytes(png)
        written.append(f"{outdir / 'figures'}/ ({len(figures)} PNG)")
    if "tsv" in formats or "xlsx" in formats:
        report = d.seqs.drop(columns=["type_label", "virus"], errors="ignore").copy()
        report["date"] = report["date"].dt.strftime("%Y-%m-%d")
        tables = {b.key: b.df for b in blocks if isinstance(b, Table)}
        if "tsv" in formats:
            (outdir / "tables").mkdir(exist_ok=True)
            report.to_csv(outdir / "report.tsv", sep="\t", index=False)
            for k, df in tables.items():
                df.to_csv(outdir / "tables" / f"{k}.tsv", sep="\t", index=False)
            written += [str(outdir / "report.tsv"), f"{outdir / 'tables'}/ ({len(tables)} TSV)"]
        if "xlsx" in formats:
            with pd.ExcelWriter(outdir / "report.xlsx", engine="openpyxl") as xw:
                report.to_excel(xw, sheet_name="sequences", index=False)
                for k, df in tables.items():
                    if k == "all_sequences":
                        continue
                    df.to_excel(xw, sheet_name=k[:31], index=False)
            written.append(str(outdir / "report.xlsx"))
    html_path = pptx_path = None
    if "html" in formats:
        html_path = outdir / "surveillance_report.html"
        html_path.write_text(render_html(d, sections), encoding="utf-8")
        written.append(str(html_path))
    if "pptx" in formats:
        pptx_path = outdir / "surveillance_report.pptx"
        render_pptx(d, sections, pptx_path, figures)
        written.append(str(pptx_path))
    verify_attribution(html_path, pptx_path)
    return written


@app.function
def main(argv=None):
    args = parse_args(argv)
    d = collect(args)
    sections = build_sections(d, strict=args.strict)
    formats = {f.strip() for f in args.formats.split(",") if f.strip()}
    written = write_outputs(d, sections, d.meta["paths"]["outdir"], formats)
    print("Surveillance report written:")
    for w in written:
        print(f"  {w}")
    return d, sections


# ======================================================================
# marimo app: live exploration (marimo run / marimo edit)
# ======================================================================


@app.function
def to_marimo(b):
    if isinstance(b, Text):
        if b.kind == "text":
            return mo.md(b.body)
        return mo.callout(mo.md(b.body), kind={"warning": "warn", "important": "danger",
                                               "error": "danger"}.get(b.kind, "info"))
    if isinstance(b, KPIs):
        return mo.vstack([mo.md(f"### {b.title}"),
                          mo.hstack([mo.stat(value=str(v), label=l, caption=c, bordered=True)
                                     for l, v, c in b.items], wrap=True, justify="start")])
    if isinstance(b, Compare):
        return mo.vstack([mo.hstack([
            mo.vstack([mo.md(f"**{h}**  \n{sub}")] + [mo.stat(value=str(v), label=l, caption=c, bordered=True)
                                                     for l, v, c in items])
            for h, sub, items in b.panels], widths="equal", align="start"), mo.md(b.caption)])
    if isinstance(b, Chart):
        return mo.vstack([mo.md(f"### {b.title}"), mo.md(b.caption),
                          mo.ui.altair_chart(b.build("container"))])
    if isinstance(b, Table):
        return mo.vstack([mo.md(f"### {b.title}"), mo.md(b.caption) if b.caption else mo.md(""),
                          mo.ui.table(b.df, page_size=b.page_length, selection=None)])
    if isinstance(b, Tabs):
        return mo.ui.tabs({label: mo.vstack([to_marimo(x) for x in inner]) for label, inner in b.tabs})
    if isinstance(b, Choice):
        return mo.ui.tabs({label: mo.vstack([to_marimo(x) for x in inner]) for label, inner in b.options})
    return mo.md("")


@app.cell(hide_code=True)
def _():
    args = parse_args(sys.argv[1:])
    is_script = mo.app_meta().mode == "script"
    return args, is_script


@app.cell(hide_code=True)
def _(args, is_script):
    # Script mode: build the report files and stop.
    if is_script:
        main(sys.argv[1:])
    return


@app.cell(hide_code=True)
def _(args, is_script):
    mo.stop(is_script)
    data = collect(args)
    as_of = mo.ui.date(value=data.meta["as_of"].date(), label="Report date (sets the season)")
    start = mo.ui.date(value=pd.Timestamp(data.cfg["plot_start_date"]).date(), label="Plot from")
    mo.vstack([
        mo.md(f"# {data.cfg['title']}\n{data.cfg.get('subtitle') or ''}"),
        mo.hstack([as_of, start], justify="start", gap=2),
    ])
    return as_of, data, start


@app.cell(hide_code=True)
def _(as_of, data, start):
    import copy
    d = copy.copy(data)
    d.cfg = dict(data.cfg, plot_start_date=str(start.value))
    d.meta = dict(data.meta, as_of=pd.Timestamp(as_of.value),
                  season_start=season_start_for(pd.Timestamp(as_of.value)))
    sections = build_sections(d)
    mo.ui.tabs({s.title: mo.vstack([mo.md(s.intro)] + [to_marimo(b) for b in s.blocks])
                for s in sections})
    return d, sections


@app.cell(hide_code=True)
def _(d):
    # Run explorer: pick a sequencing run and see every specimen's journey.
    run_picker = mo.ui.dropdown(options=list(reversed(d.run_order)), value=d.run_order[-1] if d.run_order else None,
                                label="Sequencing run", searchable=True)
    run_picker
    return (run_picker,)


@app.cell(hide_code=True)
def _(d, run_picker):
    _r = d.samples[(d.samples["sequencing_run"] == run_picker.value)]
    _cols = (["sequence_ID", "type_label", "subtype", "outcome", "outcome_detail", "subclade"]
             + [f"{s}_status" for s in BUILDS])
    mo.vstack([
        mo.hstack([mo.stat(value=str(len(_r)), label="Specimens"),
                   mo.stat(value=str(int((_r["outcome"] == "In final build").sum())), label="HA in final build"),
                   mo.stat(value=str(int((_r["genome_status"] == "in_build").sum())), label="Genome in final build")],
                  justify="start"),
        mo.ui.table(_r[[c for c in _cols if c in _r.columns]].reset_index(drop=True), page_size=25,
                    selection=None, show_column_summaries=False),
    ])
    return


@app.cell(hide_code=True)
def _(d, sections):
    export = mo.ui.run_button(label="Write HTML + PPTX to the output folder")
    export
    return (export,)


@app.cell(hide_code=True)
def _(d, export, sections):
    mo.stop(not export.value)
    _written = write_outputs(d, sections, d.meta["paths"]["outdir"], {"html", "pptx", "png", "tsv", "xlsx"})
    mo.md("Written:\n\n" + "\n".join(f"- `{w}`" for w in _written))
    return


if __name__ == "__main__":
    app.run()
